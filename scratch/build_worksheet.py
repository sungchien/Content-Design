import os
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_worksheet(output_path):
    doc = docx.Document()

    # 1. Page Margins Setup (2 cm all around for spacious, balanced layout)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    # 2. Color Palette
    COLOR_PRIMARY = RGBColor(27, 54, 93)       # Deep Navy #1B365D
    COLOR_ACCENT = RGBColor(13, 92, 117)       # Deep Teal #0D5C75
    COLOR_STAGE1 = RGBColor(21, 67, 96)        # Stage 1 Dark Blue #154360
    COLOR_STAGE2 = RGBColor(40, 55, 71)        # Stage 2 Dark Slate #283747
    COLOR_TEXT_MAIN = RGBColor(40, 40, 40)     # Charcoal #282828
    COLOR_TEXT_MUTED = RGBColor(100, 100, 100) # Gray #646464

    HEX_BG_LABEL = "F4F6F9"
    HEX_BG_WHITE = "FFFFFF"
    HEX_BORDER = "D0D7DE"
    HEX_BORDER_DARK = "1B365D"
    HEX_STAGE1_BG = "EBF5FB"
    HEX_STAGE2_BG = "EAECEE"
    HEX_CALLOUT_BG = "F8F9FA"

    # Helpers
    def set_font(run, font_name="微軟正黑體", ascii_font="Arial", size_pt=10, bold=False, italic=False, color=COLOR_TEXT_MAIN):
        run.font.name = ascii_font
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        rPr = run._r.get_or_add_rPr()
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{ascii_font}" w:hAnsi="{ascii_font}" w:eastAsia="{font_name}" w:cs="{ascii_font}"/>')
        rPr.append(rFonts)

    def set_cell_properties(cell, bg_hex=None, top_pad=100, bot_pad=100, left_pad=130, right_pad=130, valign="center", border_color=HEX_BORDER):
        tcPr = cell._tc.get_or_add_tcPr()
        if bg_hex:
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_hex}"/>')
            tcPr.append(shd)
        mar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top_pad}" w:type="dxa"/><w:bottom w:w="{bot_pad}" w:type="dxa"/><w:left w:w="{left_pad}" w:type="dxa"/><w:right w:w="{right_pad}" w:type="dxa"/></w:tcMar>')
        tcPr.append(mar)
        if valign:
            va = parse_xml(f'<w:vAlign {nsdecls("w")} w:val="{valign}"/>')
            tcPr.append(va)
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
                            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
                            f'<w:left w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
                            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
                            f'</w:tcBorders>')
        tcPr.append(borders)

    def set_table_layout(table, col_widths):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row in table.rows:
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = width

    # --- Header & Title ---
    p_super = doc.add_paragraph()
    p_super.paragraph_format.space_before = Pt(0)
    p_super.paragraph_format.space_after = Pt(2)
    r_super = p_super.add_run("數位內容設計（Digital Content Design）｜ PBL 專題實作導引（第五週）")
    set_font(r_super, font_name="微軟正黑體", ascii_font="Arial", size_pt=9.5, bold=True, color=COLOR_ACCENT)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run("第五週：SDGs 專題權威文獻洞察萃取與 Campaign 企劃收斂單")
    set_font(r_title, font_name="微軟正黑體", ascii_font="Arial", size_pt=16, bold=True, color=COLOR_PRIMARY)

    # Explanatory Callout Box (Table with left accent border)
    t_callout = doc.add_table(rows=1, cols=1)
    t_callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_callout = t_callout.cell(0, 0)
    c_callout.width = Cm(17.0)
    tcPr = c_callout._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{HEX_CALLOUT_BG}"/>'))
    tcPr.append(parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="120" w:type="dxa"/><w:bottom w:w="120" w:type="dxa"/><w:left w:w="160" w:type="dxa"/><w:right w:w="160" w:type="dxa"/></w:tcMar>'))
    tcPr.append(parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                          f'<w:top w:val="none"/>'
                          f'<w:bottom w:val="none"/>'
                          f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_BORDER_DARK}"/>'
                          f'<w:right w:val="none"/>'
                          f'</w:tcBorders>'))
    
    p_c1 = c_callout.paragraphs[0]
    p_c1.paragraph_format.space_before = Pt(0)
    p_c1.paragraph_format.space_after = Pt(3)
    p_c1.paragraph_format.line_spacing = 1.15
    r_c1_h = p_c1.add_run("【實作導引與雙階段任務說明】\n")
    set_font(r_c1_h, font_name="微軟正黑體", ascii_font="Arial", size_pt=9.5, bold=True, color=COLOR_PRIMARY)
    r_c1_b = p_c1.add_run(
        "本實作單用於第五週課堂工作坊與課後延伸任務。本週嚴格落實「知識接地（Grounding）」與「四層洞察轉譯階梯」方法論，分成兩大階段推進：\n"
        "❶【第一階段：課堂實作（60分鐘當堂查核評分）】：延續前四週確立之 SDGs 核心倡議主題與第 4 週主視覺海報資產，進行「三大權威文獻檢索活動」與「15分鐘全員矩陣社群聆聽衝刺」，在 Google NotebookLM 一鍵建立「黃金三角專屬知識庫」（官方年報＋學術實證＋社群心聲），於課堂結束前經授課教師／助教現場查核並簽章。\n"
        "❷【第二階段：課後小組延伸作業（下週開場 3 分鐘快閃分享）】：課後小組利用已建置的專屬智庫，調用四層階梯提示詞，依序萃取三大核心量化數據錨點（震撼型、對比型、切身型，點擊行內引註跳轉原文核驗）、剖析受眾深層心理阻抗，並收斂出 20 字以內 Campaign 行動標語與 100 字概念論述，完成階梯斷層審計，準備下週開場 3 分鐘小組快閃發表（Lightning Showcase）！"
    )
    set_font(r_c1_b, font_name="微軟正黑體", ascii_font="Arial", size_pt=9, color=COLOR_TEXT_MAIN)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # =========================================================================
    # STAGE 1: 課堂實作完成與現場查核項目
    # =========================================================================
    p_s1 = doc.add_paragraph()
    p_s1.paragraph_format.space_before = Pt(8)
    p_s1.paragraph_format.space_after = Pt(3)
    r_s1 = p_s1.add_run("【第一階段：課堂實作完成與現場查核項目（當堂評量，限時 60 分鐘）】")
    set_font(r_s1, font_name="微軟正黑體", ascii_font="Arial", size_pt=12, bold=True, color=COLOR_STAGE1)

    p_s1_sub = doc.add_paragraph()
    p_s1_sub.paragraph_format.space_before = Pt(0)
    p_s1_sub.paragraph_format.space_after = Pt(4)
    r_s1_sub = p_s1_sub.add_run("★ 課堂任務核心：雙向校準海報資產、三大權威文獻檢索、跨平台常民心聲採集、建置 NotebookLM 知識庫並完成當堂查核簽章。")
    set_font(r_s1_sub, font_name="微軟正黑體", ascii_font="Arial", size_pt=9, bold=False, italic=True, color=COLOR_ACCENT)

    # --- Module 1: 專案基本資訊與團隊分工 ---
    p_m1 = doc.add_paragraph()
    p_m1.paragraph_format.space_before = Pt(4)
    p_m1.paragraph_format.space_after = Pt(2)
    r_m1 = p_m1.add_run("▍模組一：專案基本資訊、視覺資產校準與團隊分工配置")
    set_font(r_m1, font_name="微軟正黑體", ascii_font="Arial", size_pt=10.5, bold=True, color=COLOR_PRIMARY)

    t1 = doc.add_table(rows=4, cols=2)
    set_table_layout(t1, [Cm(3.8), Cm(13.2)])

    rows_data_1 = [
        ("小組組別與專案名稱", "組別：第 ______ 組 ｜ 專案名稱：《__________________________________________________》"),
        ("選定之 SDGs 目標", "第 ______ 號核心目標：________________________（中文完整名稱）"),
        ("第四週視覺資產回流\n（畫面核心意象校準）",
         "【第 4 週主視覺海報主角、核心畫面意象與情感氛圍（約 40 字）】：\n"
         "__________________________________________________________________________________\n"
         "（註：審視海報主角與意象，鎖定今日檢索方向，尋找能為此海報背書的權威事實與生活摩擦力）"),
        ("小組成員與角色分工\n（課堂實作作戰矩陣）",
         "專案管理（PM / 知識庫整合）：______________ ｜ 官方政策研究員（部會年報）：______________\n"
         "學術實證研究員（期刊論文）：______________ ｜ 社群聆聽田野員（常民心聲）：______________")
    ]

    for idx, (lbl, val) in enumerate(rows_data_1):
        c0, c1 = t1.cell(idx, 0), t1.cell(idx, 1)
        set_cell_properties(c0, bg_hex=HEX_BG_LABEL, valign="center")
        set_cell_properties(c1, bg_hex=HEX_BG_WHITE, valign="top")
        
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_before = Pt(1)
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(lbl)
        set_font(r0, font_name="微軟正黑體", ascii_font="Arial", size_pt=9.5, bold=True, color=COLOR_PRIMARY)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(val)
        set_font(r1, font_name="微軟正黑體", ascii_font="Arial", size_pt=9.5, color=COLOR_TEXT_MAIN)

    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    # --- Module 2: 三大權威文獻檢索活動紀錄 ---
    p_m2 = doc.add_paragraph()
    p_m2.paragraph_format.space_before = Pt(4)
    p_m2.paragraph_format.space_after = Pt(2)
    r_m2 = p_m2.add_run("▍模組二：課堂三大權威文獻檢索活動紀錄（Authoritative Literature Hunting Sprint）")
    set_font(r_m2, font_name="微軟正黑體", ascii_font="Arial", size_pt=10.5, bold=True, color=COLOR_PRIMARY)

    t2 = doc.add_table(rows=3, cols=2)
    set_table_layout(t2, [Cm(3.8), Cm(13.2)])

    rows_data_2 = [
        ("文獻一：官方部會年報\n／國際智庫白皮書\n（Macro Policy Report）",
         "【檢索管道與高級指令語法】：\n"
         "管道：□ 管道一：政府部會出版品專區（GPI/主計總處/部會）  □ 管道二：UN/世界銀行/國際智庫\n"
         "調用之檢索語法：[例如：site:gov.tw filetype:pdf \"年報\" \"關鍵字\"]\n"
         "檢索式填寫：____________________________________________________________________\n"
         "【選定文獻資訊與品質快篩】：\n"
         "文獻全名：《__________________________________________________》\n"
         "發布機構：__________________ ｜ 發布年份：______ 年 ｜ 格式：□ PDF（已確認文字圖層可反白複製）\n"
         "決策者摘要／重點章節研究範圍（約 30 字）：________________________________________"),

        ("文獻二：同儕審查學術期刊\n／碩博士實證論文\n（Academic Study）",
         "【檢索管道與高級指令語法】：\n"
         "管道：□ 管道三：華藝線上圖書館  □ DOAJ  □ Google Scholar  □ 臺灣博碩士論文知識加值系統\n"
         "調用之檢索語法：[例如：site:edu.tw filetype:pdf \"知易行難\" OR \"態度與行為落差\" \"大學生\"]\n"
         "檢索式填寫：____________________________________________________________________\n"
         "【選定文獻資訊與品質快篩】：\n"
         "文獻全名：《__________________________________________________》\n"
         "發布機構／校院系所：__________________ ｜ 發布年份：______ 年 ｜ 格式：□ PDF（具文字圖層）\n"
         "核心實證結論／因果機制分析（約 30 字）：________________________________________"),

        ("社群田野：15 分鐘全員\n矩陣社群聆聽衝刺\n（Micro Voices）",
         "【全員跨平台採集矩陣（破除演算法同溫層）】：\n"
         "☑ Dcard 閒聊/省錢板（生活痛點）  ☑ Threads 脆（青年即時情緒）  ☑ PTT/網路新聞（反對派質疑）\n"
         "共筆檔案命名與格式：已建立《全組矩陣社群聆聽(Dcard/Threads/PTT)與思辨紀錄彙整.txt》\n"
         "【5 分鐘快閃同步會最具共鳴之常民真實聲音（代表性留言節錄）】：\n"
         "「__________________________________________________________________________________\n"
         "__________________________________________________________________________________」")
    ]

    for idx, (lbl, val) in enumerate(rows_data_2):
        c0, c1 = t2.cell(idx, 0), t2.cell(idx, 1)
        set_cell_properties(c0, bg_hex=HEX_BG_LABEL, valign="center")
        set_cell_properties(c1, bg_hex=HEX_BG_WHITE, valign="top")
        
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_before = Pt(1)
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(lbl)
        set_font(r0, font_name="微軟正黑體", ascii_font="Arial", size_pt=9.5, bold=True, color=COLOR_PRIMARY)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(val)
        set_font(r1, font_name="微軟正黑體", ascii_font="Arial", size_pt=9.5, color=COLOR_TEXT_MAIN)

    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    # --- Module 3: NotebookLM 專屬知識庫配置清單 ---
    p_m3 = doc.add_paragraph()
    p_m3.paragraph_format.space_before = Pt(4)
    p_m3.paragraph_format.space_after = Pt(2)
    r_m3 = p_m3.add_run("▍模組三：NotebookLM 黃金三角專屬知識庫配置清單")
    set_font(r_m3, font_name="微軟正黑體", ascii_font="Arial", size_pt=10.5, bold=True, color=COLOR_PRIMARY)

    t3 = doc.add_table(rows=4, cols=2)
    set_table_layout(t3, [Cm(3.8), Cm(13.2)])

    rows_data_3 = [
        ("【頂層】官方政策報告\n（宏觀法規與總量）", "文件名稱：《__________________________________________________》\n來源機構：____________________ ｜ 發布年份：______ 年 ｜ 檔案格式：PDF"),
        ("【中層】學術實證研究\n（微觀機制與落差）", "文件名稱：《__________________________________________________》\n來源機構/期刊：____________________ ｜ 發布年份：______ 年 ｜ 檔案格式：PDF"),
        ("【底層】社群田野共筆\n（常民生活摩擦力）", "文件名稱：《全組矩陣社群聆聽(Dcard/Threads/PTT)與思辨紀錄彙整.txt》\n整合內容：包含第 1-2 週 Gemini 蘇格拉底思辨紀錄 ＋ 本週跨平台熱門貼文與最多讚留言"),
        ("知識庫健康度檢查\n（來源索引確認）", "□ 3 份核心文獻已全數上傳至 NotebookLM 專屬筆記本\n□ 來源清單（Sources）全部解析成功，文字圖層無損毀或無法索引之警告標示")
    ]

    for idx, (lbl, val) in enumerate(rows_data_3):
        c0, c1 = t3.cell(idx, 0), t3.cell(idx, 1)
        set_cell_properties(c0, bg_hex=HEX_BG_LABEL, valign="center")
        set_cell_properties(c1, bg_hex=HEX_BG_WHITE, valign="top")
        
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_before = Pt(1)
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(lbl)
        set_font(r0, font_name="微軟正黑體", ascii_font="Arial", size_pt=9.5, bold=True, color=COLOR_PRIMARY)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(val)
        set_font(r1, font_name="微軟正黑體", ascii_font="Arial", size_pt=9.5, color=COLOR_TEXT_MAIN)

    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    # --- Module 4: 課堂即時諮詢與第一階段現場查核簽章 ---
    p_m4 = doc.add_paragraph()
    p_m4.paragraph_format.space_before = Pt(4)
    p_m4.paragraph_format.space_after = Pt(2)
    r_m4 = p_m4.add_run("▍模組四：課堂即時諮詢與第一階段現場查核簽章欄（當堂評分，教師／助教專用）")
    set_font(r_m4, font_name="微軟正黑體", ascii_font="Arial", size_pt=10.5, bold=True, color=COLOR_PRIMARY)

    t4 = doc.add_table(rows=2, cols=2)
    set_table_layout(t4, [Cm(3.8), Cm(13.2)])

    rows_data_4 = [
        ("第一階段查核指標\n與判定（現場勾選）",
         "【授課教師／助教現場查核判定】：\n"
         "□ 【第一階段查核通過】（符合全部三項標準，准予進入課後轉譯階梯）\n"
         "   - 指標一：文獻具備權威性（引用官方年報/學術論文，拒絕內容農場）\n"
         "   - 指標二：社群心聲具備生活摩擦力（跨平台真實聲音，反映常民痛點）\n"
         "   - 指標三：知識庫配置完整且索引正常（達成黃金三角配置，無文字圖層損毀）\n"
         "□ 【需補強後複審】（需補強項目：□ 文獻公信力不足  □ 缺少文字圖層  □ 社群聲音過於空泛）"),

        ("指導回饋意見與\n查核簽名認證",
         "【教師／助教即時諮詢指導回饋與優化建議】：\n"
         "__________________________________________________________________________________\n"
         "__________________________________________________________________________________\n"
         "查核人員簽章：______________________ ｜ 查核時間：______ 年 ______ 月 ______ 日 ______:______")
    ]

    for idx, (lbl, val) in enumerate(rows_data_4):
        c0, c1 = t4.cell(idx, 0), t4.cell(idx, 1)
        set_cell_properties(c0, bg_hex=HEX_STAGE1_BG, valign="center", border_color=HEX_BORDER_DARK)
        set_cell_properties(c1, bg_hex=HEX_BG_WHITE, valign="top", border_color=HEX_BORDER_DARK)
        
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_before = Pt(1)
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(lbl)
        set_font(r0, font_name="微軟正黑體", ascii_font="Arial", size_pt=9.5, bold=True, color=COLOR_STAGE1)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(val)
        set_font(r1, font_name="微軟正黑體", ascii_font="Arial", size_pt=9.5, color=COLOR_TEXT_MAIN)

    doc.add_page_break()

    # =========================================================================
    # STAGE 2: 課後小組延伸作業項目
    # =========================================================================
    p_s2 = doc.add_paragraph()
    p_s2.paragraph_format.space_before = Pt(4)
    p_s2.paragraph_format.space_after = Pt(3)
    r_s2 = p_s2.add_run("【第二階段：課後小組延伸作業項目（下週上課前完成並於課堂開場分享）】")
    set_font(r_s2, font_name="微軟正黑體", ascii_font="Arial", size_pt=12, bold=True, color=COLOR_STAGE2)

    p_s2_sub = doc.add_paragraph()
    p_s2_sub.paragraph_format.space_before = Pt(0)
    p_s2_sub.paragraph_format.space_after = Pt(4)
    r_s2_sub = p_s2_sub.add_run("★ 課後任務核心：調用 NotebookLM 階梯提示詞庫，依序推演四層轉譯階梯，收斂核心 Slogan，完成審計除錯，並準備下週 3 分鐘快閃發表。")
    set_font(r_s2_sub, font_name="微軟正黑體", ascii_font="Arial", size_pt=9, bold=False, italic=True, color=COLOR_ACCENT)

    # --- Module 5: 四層階梯第一層——三大核心量化數據錨點 ---
    p_m5 = doc.add_paragraph()
    p_m5.paragraph_format.space_before = Pt(4)
    p_m5.paragraph_format.space_after = Pt(2)
    r_m5 = p_m5.add_run("▍模組五：四層階梯第一層——三大核心量化數據錨點（Fact）與引註核驗")
    set_font(r_m5, font_name="微軟正黑體", ascii_font="Arial", size_pt=10.5, bold=True, color=COLOR_PRIMARY)

    t5 = doc.add_table(rows=3, cols=2)
    set_table_layout(t5, [Cm(3.8), Cm(13.2)])

    rows_data_5 = [
        ("錨點一：震撼型數據\n（Macro Impact Anchor）\n【激發危機感之鉤子】",
         "【功能定義】：巨觀危害規模（整體社會/環境受害總量，例如：全台每年損耗剩食達 384 萬公噸）\n"
         "【具體量化數據內容】：\n"
         "__________________________________________________________________________________\n"
         "【NotebookLM 行內引註出處（必填）】：\n"
         "文獻來源：《________________________》 ｜ 原文引註標籤：[ 1 ] ｜ 報告章節/頁碼：第 ____ 頁，第 ____ 段\n"
         "【一鍵跳轉查驗結論】：□ 已點擊 [1] 標籤跳轉原文對照前後文，確認數據數值與統計母體完全無誤"),

        ("錨點二：對比型數據\n（Behavior Gap Anchor）\n【引發反思之認知落差】",
         "【功能定義】：道德認同度 vs 生活執行率之懸殊對比（例如：91% 認同永續，卻僅 7% 實際隨身攜帶餐具）\n"
         "【具體量化數據內容】：\n"
         "__________________________________________________________________________________\n"
         "【NotebookLM 行內引註出處（必填）】：\n"
         "文獻來源：《________________________》 ｜ 原文引註標籤：[ 2 ] ｜ 報告章節/頁碼：第 ____ 頁，第 ____ 段\n"
         "【一鍵跳轉查驗結論】：□ 已點擊 [2] 標籤跳轉原文對照前後文，確認落差對比數據屬實無誤"),

        ("錨點三：切身型數據\n（Personal Relevance Anchor）\n【消除心理距離之痛感】",
         "【功能定義】：換算為個人日常生活必須承擔之代價（例如：平均每人每月無形多支出 2,400 元代價）\n"
         "【具體量化數據內容】：\n"
         "__________________________________________________________________________________\n"
         "【NotebookLM 行內引註出處（必填）】：\n"
         "文獻來源：《________________________》 ｜ 原文引註標籤：[ 3 ] ｜ 報告章節/頁碼：第 ____ 頁，第 ____ 段\n"
         "【一鍵跳轉查驗結論】：□ 已點擊 [3] 標籤跳轉原文對照前後文，確認個人化折算依據客觀可靠")
    ]

    for idx, (lbl, val) in enumerate(rows_data_5):
        c0, c1 = t5.cell(idx, 0), t5.cell(idx, 1)
        set_cell_properties(c0, bg_hex=HEX_BG_LABEL, valign="center")
        set_cell_properties(c1, bg_hex=HEX_BG_WHITE, valign="top")
        
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_before = Pt(1)
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(lbl)
        set_font(r0, font_name="微軟正黑體", ascii_font="Arial", size_pt=9.5, bold=True, color=COLOR_PRIMARY)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(val)
        set_font(r1, font_name="微軟正黑體", ascii_font="Arial", size_pt=9.5, color=COLOR_TEXT_MAIN)

    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    # --- Module 6: 四層階梯第二至三層——矛盾張力與受眾深層阻抗 ---
    p_m6 = doc.add_paragraph()
    p_m6.paragraph_format.space_before = Pt(4)
    p_m6.paragraph_format.space_after = Pt(2)
    r_m6 = p_m6.add_run("▍模組六：四層階梯第二至三層——矛盾張力（Tension）與受眾深層心理阻抗（Insight）")
    set_font(r_m6, font_name="微軟正黑體", ascii_font="Arial", size_pt=10.5, bold=True, color=COLOR_PRIMARY)

    t6 = doc.add_table(rows=3, cols=2)
    set_table_layout(t6, [Cm(3.8), Cm(13.2)])

    rows_data_6 = [
        ("第二層：挖掘矛盾張力\n（Tension）\n【表面觀察行為落差】",
         "【調用 2.2 步驟二提示詞，剖析官方道德宣導 vs 常民生活現實之衝突落差】：\n"
         "（例如：政策大力呼籲每餐自備保溫餐盒，但校園外宿族根本沒有流理台清洗，且期中期末熬夜趕報告只求生存便利）\n"
         "表面行為矛盾落差具體描述：\n"
         "__________________________________________________________________________________\n"
         "__________________________________________________________________________________"),

        ("第三層：社群血淋淋真實心聲\n（Social Voices）\n【青年未修飾真實心聲】",
         "【結合課堂共筆採集之 Dcard / Threads / PTT 最扎心熱門留言（原汁原味節錄）】：\n"
         "「__________________________________________________________________________________\n"
         "__________________________________________________________________________________」\n"
         "（留言出處平台與討論串主題：[平台名稱] _________________________________________）"),

        ("第三層：剖析深層心理阻抗\n（Deep Insight）\n【同理心洞察與內心獨白】",
         "【調用 2.2 步驟三提示詞，深入探究受眾「知易行難」背後的時間、空間、金錢或同儕壓力】：\n"
         "深層阻抗生活情境剖析（約 60 字）：\n"
         "__________________________________________________________________________________\n"
         "__________________________________________________________________________________\n"
         "【受眾令人鼻酸的內心獨白（一句話同理心總結）】：\n"
         "「不是我們不想支持永續，而是______________________________________________________\n"
         "逼得我們在生存現實與環境道德之間做出妥協。」")
    ]

    for idx, (lbl, val) in enumerate(rows_data_6):
        c0, c1 = t6.cell(idx, 0), t6.cell(idx, 1)
        set_cell_properties(c0, bg_hex=HEX_BG_LABEL, valign="center")
        set_cell_properties(c1, bg_hex=HEX_BG_WHITE, valign="top")
        
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_before = Pt(1)
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(lbl)
        set_font(r0, font_name="微軟正黑體", ascii_font="Arial", size_pt=9.5, bold=True, color=COLOR_PRIMARY)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(val)
        set_font(r1, font_name="微軟正黑體", ascii_font="Arial", size_pt=9.5, color=COLOR_TEXT_MAIN)

    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    # --- Module 7: 四層階梯第四層——收斂 Campaign 核心主張與快閃發表備忘 ---
    p_m7 = doc.add_paragraph()
    p_m7.paragraph_format.space_before = Pt(4)
    p_m7.paragraph_format.space_after = Pt(2)
    r_m7 = p_m7.add_run("▍模組七：四層階梯第四層——收斂 Campaign 核心主張、斷層審計與發表備忘")
    set_font(r_m7, font_name="微軟正黑體", ascii_font="Arial", size_pt=10.5, bold=True, color=COLOR_PRIMARY)

    t7 = doc.add_table(rows=4, cols=2)
    set_table_layout(t7, [Cm(3.8), Cm(13.2)])

    rows_data_7 = [
        ("第四層：一句話行動標語\n（Campaign Slogan / Hook）\n【20 字以內，全案靈魂】",
         "【調用 2.2 步驟四或 2.3 節 CRISP Mega-Prompt 收斂產出（嚴格限制 20 字以內，具備強烈情感共鳴）】：\n"
         "《________________________________________________________________》\n"
         "（標竿參考範例：《外表不完美，風味更甜美！月底深夜五星級快閃食堂》）"),

        ("第四層：核心傳播概念論述\n（Concept Brief）\n【約 100 字，跨通路靈魂】",
         "【全案核心傳播論述說明（約 100 字，清楚交代『數據痛點 ➔ 心理阻抗 ➔ 創意解方』三位一體之連結）】：\n"
         "__________________________________________________________________________________\n"
         "__________________________________________________________________________________\n"
         "__________________________________________________________________________________"),

        ("階梯斷層審計自我診斷\n（Ladder Audit Prompt 自檢）\n【杜絕三大典型盲區】",
         "【調用講義 2.4 節階梯斷層審計提示詞，逐項自我診斷並勾選】：\n"
         "□ 【斷層一檢驗（無資訊搬運）】：標語具備生活溫度與情感張力，絕無直接把生硬百分比抄成口號\n"
         "□ 【斷層二檢驗（無道德說教）】：論述同理受眾生活摩擦力，絕無指責受眾「不道德或浪費」引發反感\n"
         "□ 【斷層三檢驗（無口號脫節）】：倡議主張緊扣前面三大數據錨點與痛點，具備具體可落地的行動載體"),

        ("下週開場 3 分鐘快閃發表\n準備備忘錄\n（Lightning Showcase）",
         "【下週課堂開場快閃口頭發表分工與講稿骨架備忘】：\n"
         "發表代表同學：______________ ｜ 搭配展示之第 4 週海報／簡報檔案：____________________\n"
         "3 分鐘極速發表講稿架構（每部分各 1 分鐘）：\n"
         "• 1.【視覺與數據鉤子（1 分鐘）】：展示主視覺海報 ➔ 拋出震撼型或切身型數據錨點，抓住全場耳朵\n"
         "• 2.【社群真實痛點（1 分鐘）】：唸出最扎心的社群留言 ➔ 揭示受眾「知易行難」背後令人鼻酸的內心獨白\n"
         "• 3.【核心主張亮劍（1 分鐘）】：全組喊出 20 字 Campaign 口號 ➔ 闡述 100 字概念論述，接受同儕回饋")
    ]

    for idx, (lbl, val) in enumerate(rows_data_7):
        c0, c1 = t7.cell(idx, 0), t7.cell(idx, 1)
        set_cell_properties(c0, bg_hex=HEX_BG_LABEL, valign="center")
        set_cell_properties(c1, bg_hex=HEX_BG_WHITE, valign="top")
        
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_before = Pt(1)
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(lbl)
        set_font(r0, font_name="微軟正黑體", ascii_font="Arial", size_pt=9.5, bold=True, color=COLOR_PRIMARY)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(val)
        set_font(r1, font_name="微軟正黑體", ascii_font="Arial", size_pt=9.5, color=COLOR_TEXT_MAIN)

    # Save
    doc.save(output_path)
    print(f"Successfully generated: {output_path}")

if __name__ == "__main__":
    output_file = os.path.join(os.getcwd(), "SDGs_專題權威文獻洞察萃取與Campaign企劃收斂單.docx")
    create_worksheet(output_file)
