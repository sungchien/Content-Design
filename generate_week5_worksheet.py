import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = docx.Document()
    
    # 頁面邊界設定 (2.032 cm / 0.8 inch 全等邊界，維持現有 Word 實作單規格)
    section = doc.sections[0]
    section.top_margin = Cm(2.032)
    section.bottom_margin = Cm(2.032)
    section.left_margin = Cm(2.032)
    section.right_margin = Cm(2.032)
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    
    # 全域樣式設定
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Microsoft JhengHei'
    normal_style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    normal_style.font.size = Pt(9.5)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(3)

    # 樣式常數定義
    COLOR_PRIMARY = RGBColor(0x1a, 0x56, 0xb8)     # 皇家藍 (標題與章節標題)
    COLOR_MUTED = RGBColor(0x7f, 0x8c, 0x8d)       # 灰字 (頂部引導備註)
    COLOR_BODY = RGBColor(0x55, 0x55, 0x55)        # 說明次要字
    COLOR_BENCHMARK = RGBColor(0x2e, 0x7d, 0x32)   # 墨綠色 (示範標竿)
    COLOR_REVIEW = RGBColor(0xc0, 0x39, 0x2b)      # 緋紅色 (教師審查欄)
    
    FILL_HEADER = "F0F4F8"                         # 欄位標籤淺藍灰底色
    BORDER_COLOR = "D0D7DE"                        # 表格細邊框淺灰色
    
    TOTAL_WIDTH_DXA = 9936                         # 17.526 cm (可列印區域總寬度)
    LABEL_WIDTH_DXA = 2880                         # 5.08 cm
    CONTENT_WIDTH_DXA = 7056                       # 12.446 cm
    
    def set_cell_format(cell, width_dxa=None, fill_hex=None, v_align="center", pad_top=100, pad_bottom=100, pad_left=130, pad_right=130):
        tcPr = cell._tc.get_or_add_tcPr()
        
        # 寬度
        if width_dxa is not None:
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW')
                tcPr.append(tcW)
            tcW.set(qn('w:type'), 'dxa')
            tcW.set(qn('w:w'), str(width_dxa))
            
        # 垂直對齊
        vAlign = tcPr.find(qn('w:vAlign'))
        if vAlign is None:
            vAlign = OxmlElement('w:vAlign')
            tcPr.append(vAlign)
        vAlign.set(qn('w:val'), v_align)
        
        # 內距 padding
        tcMar = tcPr.find(qn('w:tcMar'))
        if tcMar is None:
            tcMar = OxmlElement('w:tcMar')
            tcPr.append(tcMar)
        for side, pad in [('top', pad_top), ('bottom', pad_bottom), ('left', pad_left), ('right', pad_right)]:
            m = tcMar.find(qn(f'w:{side}'))
            if m is None:
                m = OxmlElement(f'w:{side}')
                tcMar.append(m)
            m.set(qn('w:w'), str(pad))
            m.set(qn('w:type'), 'dxa')
            
        # 邊框
        borders_xml = f'''
        <w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>
        </w:tcBorders>
        '''
        existing_borders = tcPr.find(qn('w:tcBorders'))
        if existing_borders is not None:
            tcPr.remove(existing_borders)
        tcPr.append(parse_xml(borders_xml))
        
        # 底色
        if fill_hex:
            shd = tcPr.find(qn('w:shd'))
            if shd is None:
                shd = OxmlElement('w:shd')
                tcPr.append(shd)
            shd.set(qn('w:fill'), fill_hex)

    def set_table_props(table):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = table._tbl.tblPr
        tblW = tblPr.find(qn('w:tblW'))
        if tblW is not None:
            tblW.set(qn('w:type'), 'auto')
            tblW.set(qn('w:w'), '0')
        for row in table.rows:
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml('<w:cantSplit xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))

    def add_heading_paragraph(text, color=COLOR_PRIMARY, size_pt=11.5, space_before=10, space_after=3):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Microsoft JhengHei'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
        run.font.size = Pt(size_pt)
        run.font.bold = True
        run.font.color.rgb = color
        return p

    # ------------------ 文件頂部抬頭 ------------------
    p0 = doc.add_paragraph()
    p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p0.paragraph_format.space_after = Pt(2)
    r0 = p0.add_run('數位內容設計（Digital Content Design）｜ PBL 專題實作導引（第五週）')
    r0.font.name = 'Microsoft JhengHei'
    r0._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    r0.font.size = Pt(9.0)
    r0.font.color.rgb = COLOR_MUTED

    p1 = doc.add_paragraph()
    p1.paragraph_format.space_after = Pt(4)
    p1.paragraph_format.space_before = Pt(0)
    r1 = p1.add_run('第五週：SDGs 專題權威文獻洞察萃取與 Campaign 企劃收斂單')
    r1.font.name = 'Microsoft JhengHei'
    r1._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    r1.font.size = Pt(17.0)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_PRIMARY

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(8)
    p2.paragraph_format.line_spacing = 1.25
    r2 = p2.add_run(
        '本實作單用於第五週課堂工作坊（60 分鐘）。請各組延續前四週所確立之 SDGs 核心倡議主題與主視覺資產，實踐「知識接地（Grounding）與文獻洞察轉譯工作流」：'
        '首先於 Google NotebookLM（notebooklm.google.com）匯入至少 2 份聯合國或政府官方發布之權威白皮書與實證學術文獻，建構小組專屬知識庫；'
        '接著運用結構化文獻提問法，萃取三大核心量化數據錨點（震撼型、對比型、切身型），並透過行內引註（In-line Citations）逐一跳轉原文核驗；'
        '隨後探討目標受眾「知易行難」背後之深層心理阻抗，將數據與人性痛點收斂為一組具備高度傳播爆發力的一句話 Campaign Slogan 與 100 字概念論述；'
        '最後試聽 Audio Overview 深度音訊總覽，提煉生動的生活化比喻，作為後續影音腳本之敘事靈感。'
        '完成後於課堂結束前提交，經教師/助教諮詢查核後，作為專案期末企劃書與活動網頁之核心論據資產。'
    )
    r2.font.name = 'Microsoft JhengHei'
    r2._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    r2.font.size = Pt(9.0)
    r2.font.color.rgb = COLOR_BODY

    # ------------------ 表格一：基本資訊與團隊分工 ------------------
    add_heading_paragraph('【模組一：專案基本資訊與團隊分工配置】')
    
    t1 = doc.add_table(rows=4, cols=2)
    set_table_props(t1)
    
    fields_t1 = [
        ('小組組別與專案名稱', '組別：第 ______ 組 ｜ 專案名稱：《__________________________________________________》'),
        ('選定之 SDGs 目標', '第 ______ 號核心目標：________________________（中文完整名稱）'),
        ('第四週視覺資產回顧\n（畫面核心意象）', '【主視覺海報主角與核心畫面意象（約 40 字）】：\n__________________________________________________________________________________'),
        ('小組成員與角色分工', '專案管理（PM）：______________ ｜ 內容策略與文案：______________\n視覺設計與品牌：______________ ｜ 多模態影音製作：______________')
    ]
    
    for i, (label, val) in enumerate(fields_t1):
        c0 = t1.cell(i, 0)
        c1 = t1.cell(i, 1)
        set_cell_format(c0, width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER)
        set_cell_format(c1, width_dxa=CONTENT_WIDTH_DXA)
        
        p_lbl = c0.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(9.5)
        
        p_val = c1.paragraphs[0]
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.5)

    # ------------------ 表格二：NotebookLM 知識庫來源建置 ------------------
    add_heading_paragraph('【模組二：NotebookLM 專屬知識庫文獻建置紀錄】')
    
    t2 = doc.add_table(rows=3, cols=2)
    set_table_props(t2)
    
    fields_t2 = [
        ('官方權威報告一\n（Macro Policy Report）', '報告名稱：《__________________________________________________》\n發布機構：__________________ ｜ 發布年份：______ 年 ｜ 格式：□ PDF  □ 網頁 URL\n核心研究範圍（約 30 字）：_____________________________________________________'),
        ('官方報告或學術文獻二\n（Academic / Policy Study）', '報告名稱：《__________________________________________________》\n發布機構：__________________ ｜ 發布年份：______ 年 ｜ 格式：□ PDF  □ 網頁 URL\n核心研究範圍（約 30 字）：_____________________________________________________'),
        ('在地田野或社群資料三\n（Micro Voices / 選填）', '資料名稱：《__________________________________________________》\n來源類型：□ NGO 訪談筆記  □ 社群討論串整理  □ 第二週蘇格拉底思辨表\n核心內容摘要（約 30 字）：_____________________________________________________')
    ]
    
    for i, (label, val) in enumerate(fields_t2):
        c0 = t2.cell(i, 0)
        c1 = t2.cell(i, 1)
        set_cell_format(c0, width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER)
        set_cell_format(c1, width_dxa=CONTENT_WIDTH_DXA)
        
        p_lbl = c0.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(9.5)
        
        p_val = c1.paragraphs[0]
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.0)

    # ------------------ 表格三：三大量化數據錨點提煉 ------------------
    add_heading_paragraph('【模組三：三大核心量化數據錨點（Data Anchors）與引註溯源】')
    
    t3 = doc.add_table(rows=3, cols=2)
    set_table_props(t3)
    
    fields_t3 = [
        ('錨點一：震撼型數據\n（巨觀危害規模）\n【用於吸引眼球之鉤子】', '【量化數據內容】（例如：每年產生損耗公噸數、全球受害總量）：\n__________________________________________________________________________________\n【原文引註出處（必填）】：\n文獻來源：________________________ ｜ 報告章節/頁碼：第 ____ 頁，第 ____ 段\n【一鍵跳轉查驗結論】：□ 已點擊核對原文前後文，確認無誤'),
        ('錨點二：對比型數據\n（行動認知落差）\n【用於引發社群思辨】', '【量化數據內容】（例如：80% 認同永續 vs 僅 5% 真正自備餐具）：\n__________________________________________________________________________________\n【原文引註出處（必填）】：\n文獻來源：________________________ ｜ 報告章節/頁碼：第 ____ 頁，第 ____ 段\n【一鍵跳轉查驗結論】：□ 已點擊核對原文前後文，確認無誤'),
        ('錨點三：切身型數據\n（生活日常關聯）\n【用於消除心理距離】', '【量化數據內容】（例如：平均每位大學生每年多花 3,000 元承擔無形成本）：\n__________________________________________________________________________________\n【原文引註出處（必填）】：\n文獻來源：________________________ ｜ 報告章節/頁碼：第 ____ 頁，第 ____ 段\n【一鍵跳轉查驗結論】：□ 已點擊核對原文前後文，確認無誤')
    ]
    
    for i, (label, val) in enumerate(fields_t3):
        c0 = t3.cell(i, 0)
        c1 = t3.cell(i, 1)
        set_cell_format(c0, width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER)
        set_cell_format(c1, width_dxa=CONTENT_WIDTH_DXA)
        
        p_lbl = c0.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(9.5)
        
        p_val = c1.paragraphs[0]
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.0)

    # ------------------ 表格四：質化洞察與 Campaign 概念收斂 ------------------
    add_heading_paragraph('【模組四：受眾深層阻抗剖析與核心 Campaign 主張收斂】')
    
    t4 = doc.add_table(rows=3, cols=2)
    set_table_props(t4)
    
    fields_t4 = [
        ('受眾深層心理阻抗\n（質化同理心洞察）', '【目標受眾表面行為落差】：\n__________________________________________________________________________________\n【知易行難背後的真實困境與心理阻抗（約 60 字）】：\n（例如：並非缺乏環保意識，而是月底經濟拮据、缺乏料理器具，生存便利性被迫壓過道德感）\n__________________________________________________________________________________'),
        ('核心行動標語\n（Campaign Slogan / Hook）\n【20 字以內，直擊痛點】', '《________________________________________________________________》\n（標竿參考：《外表不完美，風味更甜美！月底深夜五星級快閃食堂》）'),
        ('核心傳播概念論述\n（Campaign Brief）\n【約 100 字，全案靈魂】', '【全案傳播核心主張說明（約 100 字）】：\n__________________________________________________________________________________\n__________________________________________________________________________________\n__________________________________________________________________________________')
    ]
    
    for i, (label, val) in enumerate(fields_t4):
        c0 = t4.cell(i, 0)
        c1 = t4.cell(i, 1)
        set_cell_format(c0, width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER)
        set_cell_format(c1, width_dxa=CONTENT_WIDTH_DXA)
        
        p_lbl = c0.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(9.5)
        
        p_val = c1.paragraphs[0]
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.0)

    # ------------------ 表格五：Audio Overview 轉譯與審查欄 ------------------
    add_heading_paragraph('【模組五：Audio Overview 多模態轉譯反思與課堂審查】')
    
    t5 = doc.add_table(rows=2, cols=2)
    set_table_props(t5)
    
    fields_t5 = [
        ('Audio Overview\n語音播客轉譯反思\n（多模態敘事靈感）', '【AI 雙人對談中出現之精彩生活化比喻或衝突對白（約 60～80 字）】：\n（例如：主持人將微塑膠比喻為每週吞下一張信用卡，或反思洗衣機排出的合成纖維才是隱形殺手）\n__________________________________________________________________________________\n【預計如何轉化為宣傳短片之開場旁白？】：\n__________________________________________________________________________________'),
        ('課堂即時諮詢與\n審查回饋欄\n（授課教師／助教專用）', '【專案立論與文獻接地審查判定】：\n□ 審核通過（數據扎實、引註精確、主張直擊痛點）\n□ 建議微調（數據權威性不足或單位缺乏生活體感）\n□ 重新收斂（口號與文獻數據嚴重脫節）\n\n指導回饋與具體修改意見：\n__________________________________________________________________________________\n\n教師／助教簽章：______________________ ｜ 審核日期：2026 年 10 月 14 日')
    ]
    
    for i, (label, val) in enumerate(fields_t5):
        c0 = t5.cell(i, 0)
        c1 = t5.cell(i, 1)
        if i == 1:
            set_cell_format(c0, width_dxa=LABEL_WIDTH_DXA, fill_hex="FDEDEC")
            set_cell_format(c1, width_dxa=CONTENT_WIDTH_DXA, fill_hex="FDFEFE")
            p_lbl = c0.paragraphs[0]
            r_lbl = p_lbl.add_run(label)
            r_lbl.font.bold = True
            r_lbl.font.size = Pt(9.5)
            r_lbl.font.color.rgb = COLOR_REVIEW
        else:
            set_cell_format(c0, width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER)
            set_cell_format(c1, width_dxa=CONTENT_WIDTH_DXA)
            p_lbl = c0.paragraphs[0]
            r_lbl = p_lbl.add_run(label)
            r_lbl.font.bold = True
            r_lbl.font.size = Pt(9.5)
        
        p_val = c1.paragraphs[0]
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.0)

    # 存檔
    output_path = 'SDGs_專題權威文獻洞察萃取與Campaign企劃收斂單.docx'
    doc.save(output_path)
    print("Successfully generated worksheet docx file.")

if __name__ == '__main__':
    create_document()
