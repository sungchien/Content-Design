import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = docx.Document()
    
    # 頁面邊界設定 (2.032 cm / 0.8 inch 全等邊界，維持現有 Word 實作單規格)
    section = doc.sections[0]
    section.top_margin = Cm(2.032)
    section.bottom_margin = Cm(2.032)
    section.left_margin = Cm(2.032)
    section.right_margin = Cm(2.032)
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    
    # 全域樣式設定
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Microsoft JhengHei'
    normal_style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    normal_style.font.size = Pt(9.5)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(3)

    # 樣式常數定義
    COLOR_PRIMARY = RGBColor(0x1a, 0x56, 0xb8)     # 皇家藍 (標題與章節標題)
    COLOR_MUTED = RGBColor(0x7f, 0x8c, 0x8d)       # 灰字 (頂部引導備註)
    COLOR_BODY = RGBColor(0x55, 0x55, 0x55)        # 說明次要字
    COLOR_BENCHMARK = RGBColor(0x2e, 0x7d, 0x32)   # 墨綠色 (示範標竿)
    COLOR_REVIEW = RGBColor(0xc0, 0x39, 0x2b)      # 緋紅色 (教師審查欄)
    
    FILL_HEADER = "F0F4F8"                         # 欄位標籤淺藍灰底色
    BORDER_COLOR = "D0D7DE"                        # 表格細邊框淺灰色
    
    TOTAL_WIDTH_DXA = 9936                         # 17.526 cm (可列印區域總寬度)
    LABEL_WIDTH_DXA = 2880                         # 5.08 cm
    CONTENT_WIDTH_DXA = 7056                       # 12.446 cm
    
    def set_cell_format(cell, width_dxa=None, fill_hex=None, v_align="center", pad_top=100, pad_bottom=100, pad_left=130, pad_right=130):
        tcPr = cell._tc.get_or_add_tcPr()
        
        # 寬度
        if width_dxa is not None:
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW')
                tcPr.append(tcW)
            tcW.set(qn('w:type'), 'dxa')
            tcW.set(qn('w:w'), str(width_dxa))
            
        # 垂直對齊
        vAlign = tcPr.find(qn('w:vAlign'))
        if vAlign is None:
            vAlign = OxmlElement('w:vAlign')
            tcPr.append(vAlign)
        vAlign.set(qn('w:val'), v_align)
        
        # 內距 padding
        tcMar = tcPr.find(qn('w:tcMar'))
        if tcMar is None:
            tcMar = OxmlElement('w:tcMar')
            tcPr.append(tcMar)
        for side, pad in [('top', pad_top), ('bottom', pad_bottom), ('left', pad_left), ('right', pad_right)]:
            m = tcMar.find(qn(f'w:{side}'))
            if m is None:
                m = OxmlElement(f'w:{side}')
                tcMar.append(m)
            m.set(qn('w:w'), str(pad))
            m.set(qn('w:type'), 'dxa')
            
        # 邊框
        borders_xml = f'''
        <w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>
        </w:tcBorders>
        '''
        existing_borders = tcPr.find(qn('w:tcBorders'))
        if existing_borders is not None:
            tcPr.remove(existing_borders)
        tcPr.append(parse_xml(borders_xml))
        
        # 底色
        if fill_hex:
            shd = tcPr.find(qn('w:shd'))
            if shd is None:
                shd = OxmlElement('w:shd')
                tcPr.append(shd)
            shd.set(qn('w:fill'), fill_hex)

    def set_table_props(table):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = table._tbl.tblPr
        tblW = tblPr.find(qn('w:tblW'))
        if tblW is not None:
            tblW.set(qn('w:type'), 'auto')
            tblW.set(qn('w:w'), '0')
        # 防止跨頁截斷行
        for row in table.rows:
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml('<w:cantSplit xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))

    def add_heading_paragraph(text, color=COLOR_PRIMARY, size_pt=11.5, space_before=10, space_after=3):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Microsoft JhengHei'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
        run.font.size = Pt(size_pt)
        run.font.bold = True
        run.font.color.rgb = color
        return p

    # ------------------ 文件頂部抬頭 ------------------
    p0 = doc.add_paragraph()
    p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p0.paragraph_format.space_after = Pt(2)
    r0 = p0.add_run('數位內容設計（Digital Content Design）｜ PBL 專題實作導引（第三週）')
    r0.font.name = 'Microsoft JhengHei'
    r0._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    r0.font.size = Pt(9.0)
    r0.font.color.rgb = COLOR_MUTED

    p1 = doc.add_paragraph()
    p1.paragraph_format.space_after = Pt(4)
    p1.paragraph_format.space_before = Pt(0)
    r1 = p1.add_run('第三週：SDGs 結構化提示工程與多情境資料擴展實作單')
    r1.font.name = 'Microsoft JhengHei'
    r1._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    r1.font.size = Pt(17.0)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_PRIMARY

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(8)
    p2.paragraph_format.line_spacing = 1.25
    r2 = p2.add_run(
        '本實作單用於第三週課堂工作坊（60 分鐘）。請各組打開 Google AI Studio（aistudio.google.com），'
        '延續第二週蘇格拉底思辨所淬鍊之「真痛點」與「利害關係人抗拒洞察」，建立小組專屬之 CRISP 結構化提示詞架構，'
        '並完成多面向受眾共感文案擴展與 Few-shot 反對聲浪意圖分類測試。完成後於課堂結束前提交，經教師/助教諮詢查核後，'
        '作為後續第四週多模態生成（Imagen / Nano Banana）與專案內容行銷之核心提示詞引擎。'
    )
    r2.font.name = 'Microsoft JhengHei'
    r2._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = COLOR_BODY

    # ------------------ 一、 小組基本資訊與 AI Studio 提示工程分工 ------------------
    add_heading_paragraph('一、 小組基本資訊與 AI Studio 提示工程分工')
    t0 = doc.add_table(rows=3, cols=4)
    set_table_props(t0)
    col_w_t0 = 2484  # 9936 / 4 = 2484 dxa
    
    # R0
    t0.cell(0, 0).paragraphs[0].add_run('專案組別').bold = True
    t0.cell(0, 1).paragraphs[0].add_run('第 _______ 組')
    t0.cell(0, 2).paragraphs[0].add_run('小組組名').bold = True
    t0.cell(0, 3).paragraphs[0].add_run('《 ___________________ 》')
    set_cell_format(t0.cell(0, 0), width_dxa=col_w_t0, fill_hex=FILL_HEADER)
    set_cell_format(t0.cell(0, 1), width_dxa=col_w_t0)
    set_cell_format(t0.cell(0, 2), width_dxa=col_w_t0, fill_hex=FILL_HEADER)
    set_cell_format(t0.cell(0, 3), width_dxa=col_w_t0)

    # R1
    t0.cell(1, 0).paragraphs[0].add_run('專案暫定名稱').bold = True
    t0.cell(1, 1).paragraphs[0].add_run('《 _______________________________ 》')
    t0.cell(1, 2).paragraphs[0].add_run('選定 SDGs 目標').bold = True
    t0.cell(1, 3).paragraphs[0].add_run('第 【    】 號目標：_________________')
    set_cell_format(t0.cell(1, 0), width_dxa=col_w_t0, fill_hex=FILL_HEADER)
    set_cell_format(t0.cell(1, 1), width_dxa=col_w_t0)
    set_cell_format(t0.cell(1, 2), width_dxa=col_w_t0, fill_hex=FILL_HEADER)
    set_cell_format(t0.cell(1, 3), width_dxa=col_w_t0)

    # R2 (合併後 3 格)
    t0.cell(2, 0).paragraphs[0].add_run('今日提示工程\n職能分工').bold = True
    set_cell_format(t0.cell(2, 0), width_dxa=col_w_t0, fill_hex=FILL_HEADER)
    
    c_merged = t0.cell(2, 1).merge(t0.cell(2, 3))
    p_roles = c_merged.paragraphs[0]
    p_roles.paragraph_format.line_spacing = 1.3
    r_role1 = p_roles.add_run('• 提示架構師 (Prompt Engineer / PM)：____________________（負責統籌 CRISP 五大構件、System Instructions 與參數調校）\n')
    r_role2 = p_roles.add_run('• 文案策略員 (Copy Strategist / Content)：________________（負責次受眾痛點切角轉譯與三版本社群共感文案調校）\n')
    r_role3 = p_roles.add_run('• 資料標記員 (Data Classifier / Creative)：_______________（負責收集真實反對聲浪情境並建立 Few-shot 分類範例）')
    set_cell_format(c_merged, width_dxa=col_w_t0 * 3, v_align="top", pad_top=120, pad_bottom=120)

    # ------------------ 二、 Google AI Studio 環境配置與 System Instructions 全域守則 ------------------
    add_heading_paragraph('二、 Google AI Studio 環境配置與 System Instructions 全域守則')
    t1 = doc.add_table(rows=2, cols=2)
    set_table_props(t1)

    # R0: 模型與解碼參數
    t1.cell(0, 0).paragraphs[0].add_run('AI Studio 模型與\n解碼參數配置\n\n【思考轉盤調控】').bold = True
    set_cell_format(t1.cell(0, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_prm = t1.cell(0, 1).paragraphs[0]
    p_prm.paragraph_format.line_spacing = 1.35
    p_prm.add_run('1. 選用模型版本：☐ Gemini Pro Latest (Gemini 3.1 Pro)    ☐ Gemini Flash Latest\n')
    p_prm.add_run('2. 任務一（多面向受眾文案擴展，創造發散）：\n')
    p_prm.add_run('   • Temperature 設定值：________ ｜ Top-p 設定值：________\n')
    p_prm.add_run('   • 參數調校依據（請以「思考轉盤發散度」簡要說明，約 30 字）：\n     __________________________________________________________________________________\n')
    p_prm.add_run('3. 任務二（反對聲浪意圖分類，嚴謹收斂）：\n')
    p_prm.add_run('   • Temperature 設定值：________ ｜ Top-p 設定值：________\n')
    p_prm.add_run('   • 參數調校依據（請以「思考轉盤收斂度」簡要說明，約 30 字）：\n     __________________________________________________________________________________')
    set_cell_format(t1.cell(0, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=120, pad_bottom=120)

    # R1: System Instructions 全域設定
    t1.cell(1, 0).paragraphs[0].add_run('System Instructions\n全域角色設定\n\n【鎖定模型靈魂】').bold = True
    set_cell_format(t1.cell(1, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_sys = t1.cell(1, 1).paragraphs[0]
    p_sys.paragraph_format.line_spacing = 1.3
    p_sys.add_run('請填寫於 AI Studio 右側 "System Instructions" 欄位之設定（杜絕道德說教，鎖定生活化視角與台灣習慣用語）：\n').font.color.rgb = COLOR_BODY
    p_sys.add_run('你是專精於【 ______________________________________ 】的數位內容策略教練。\n')
    p_sys.add_run('你的唯一職責是協助青年團隊將複雜的永續議題轉化為具備高社群感染力的內容。\n')
    p_sys.add_run('在所有後續回應中，你必須嚴格遵守以下守則：\n')
    p_sys.add_run('1. 嚴禁使用居高臨下的道德說教口吻（如「大家必須愛護地球，否則會遭到報應」）。\n')
    p_sys.add_run('2. 一律採用生活化微觀視角，將環境成本與大專學生的日常生活（外送、網購、外宿料理）直接連結。\n')
    p_sys.add_run('3. 所有產出均以繁體中文（台灣習慣用語）呈現，禁止使用簡體中文或中國大陸慣用行銷術語。\n')
    p_sys.add_run('小組自訂守則（選填）：4. ________________________________________________________________________')
    set_cell_format(t1.cell(1, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=120, pad_bottom=120)

    # ------------------ 三、 小組專屬 CRISP 結構化提示詞架構（核心提示詞引擎） ------------------
    add_heading_paragraph('三、 小組專屬 CRISP 結構化提示詞架構（核心提示詞引擎）')
    t2 = doc.add_table(rows=5, cols=2)
    set_table_props(t2)

    crisp_items = [
        ('【C - Context】\n背景語境', 
         '提供專案背景與宏觀現況（例：校園外宿生三餐外食、外送一次性餐具激增、每週製造大量塑膠垃圾...）：\n\n\n'),
        ('【R - Role】\n專家角色', 
         '賦予 AI 具體的權威角色與語氣特徵（例：獲獎無數的永續青年社群行銷總監，擅長以犀利且幽默的自嘲風格引發大學生共鳴...）：\n\n\n'),
        ('【I - Instruction】\n核心指令', 
         '清晰下達動詞驅動的明確任務（例：請針對目標受眾撰寫 3 組具備高點擊潛力的社群貼文文案，每組包含標題、痛點內文與行動呼籲...）：\n\n\n'),
        ('【S - Scenario】\n具體場景', 
         '描述受眾的生活場景與心理狀態（例：月底荷包見底、晚上 11 點在宿舍猶豫要不要點宵夜外送的疲憊大學生...）：\n\n\n'),
        ('【P - Product】\n產出結果規格', 
         '嚴格限制格式、字數、欄位與語系（例：以 Markdown 表格呈現，包含【次受眾類型】、【Hook 標題】、【痛點訴求（80字內）】、【行動呼籲 CTA】...）：\n\n\n')
    ]

    for idx, (label, prompt_desc) in enumerate(crisp_items):
        t2.cell(idx, 0).paragraphs[0].add_run(label).bold = True
        set_cell_format(t2.cell(idx, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")
        
        p = t2.cell(idx, 1).paragraphs[0]
        p.paragraph_format.line_spacing = 1.3
        p.add_run(prompt_desc)
        set_cell_format(t2.cell(idx, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=100, pad_bottom=100)

    # ------------------ 四、 任務一：多面向受眾共感文案擴展成果（擴展性任務） ------------------
    add_heading_paragraph('四、 任務一：多面向受眾共感文案擴展成果（擴展性任務）')
    t3 = doc.add_table(rows=3, cols=2)
    set_table_props(t3)

    audience_items = [
        ('次受眾 A 文案\n【經濟 / 省錢導向】', 
         '• 次受眾生活畫像與核心痛點：____________________________________________________\n'
         '• 吸睛標題 (Hook)：____________________________________________________________\n'
         '• 核心訴求內文 (約 80 字)：\n\n\n'
         '• 行動呼籲 (Call To Action, CTA)：______________________________________________'),
        ('次受眾 B 文案\n【健康 / 生活焦慮導向】', 
         '• 次受眾生活畫像與核心痛點：____________________________________________________\n'
         '• 吸睛標題 (Hook)：____________________________________________________________\n'
         '• 核心訴求內文 (約 80 字)：\n\n\n'
         '• 行動呼籲 (Call To Action, CTA)：______________________________________________'),
        ('次受眾 C 文案\n【美學 / 儀式感導向】', 
         '• 次受眾生活畫像與核心痛點：____________________________________________________\n'
         '• 吸睛標題 (Hook)：____________________________________________________________\n'
         '• 核心訴求內文 (約 80 字)：\n\n\n'
         '• 行動呼籲 (Call To Action, CTA)：______________________________________________')
    ]

    for idx, (label, content) in enumerate(audience_items):
        t3.cell(idx, 0).paragraphs[0].add_run(label).bold = True
        set_cell_format(t3.cell(idx, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")
        
        p = t3.cell(idx, 1).paragraphs[0]
        p.paragraph_format.line_spacing = 1.3
        p.add_run(content)
        set_cell_format(t3.cell(idx, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=100, pad_bottom=100)

    # ------------------ 五、 任務二：社群反對聲浪意圖與情緒極性標記成果（Few-shot 分類任務） ------------------
    add_heading_paragraph('五、 任務二：社群反對聲浪意圖與情緒極性標記成果（Few-shot 分類任務）')
    t4 = doc.add_table(rows=3, cols=2)
    set_table_props(t4)

    t4.cell(0, 0).paragraphs[0].add_run('Few-shot 分類器\n範例設定\n\n【示範對照組】').bold = True
    set_cell_format(t4.cell(0, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_fs = t4.cell(0, 1).paragraphs[0]
    p_fs.paragraph_format.line_spacing = 1.3
    p_fs.add_run('小組於 Prompt 中提供的 Few-shot 輸入範例與期望輸出（提供模型學習標準格式）：\n').font.color.rgb = COLOR_BODY
    p_fs.add_run('【輸入範例】：「循環杯真的有洗乾淨嗎？之前借過一次裡面還有口紅印，超級噁心，打死都不會借！」\n')
    p_fs.add_run('【期望輸出】：\n')
    p_fs.add_run('  • 留言核心意圖：衛生與清潔信任疑慮\n')
    p_fs.add_run('  • 情緒極性類別：極度負面 (Strongly Negative)\n')
    p_fs.add_run('  • 受眾心理痛點：害怕傳染病菌、對公共清洗消毒流程缺乏信任度\n')
    p_fs.add_run('  • 建議化解策略：拍攝『循環杯高溫殺菌清洗廠 30 秒透明化縮時影片』，主打洗淨度超越個人手洗。')
    set_cell_format(t4.cell(0, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=110, pad_bottom=110)

    feedback_tests = [
        ('實測質疑一\n【成本 / 利益衝突】', 
         '• 輸入之真實社群質疑情境：____________________________________________________\n'
         '• 模型自動標記之情緒極性：☐ 強烈負面    ☐ 懷疑/中立    ☐ 建設性建議\n'
         '• 模型診斷之核心抗拒原因：____________________________________________________\n'
         '• 公關文案化解與溝通對策：____________________________________________________\n'
         '  __________________________________________________________________________________'),
        ('實測質疑二\n【便利性 / 嫌麻煩】', 
         '• 輸入之真實社群質疑情境：____________________________________________________\n'
         '• 模型自動標記之情緒極性：☐ 強烈負面    ☐ 懷疑/中立    ☐ 建設性建議\n'
         '• 模型診斷之核心抗拒原因：____________________________________________________\n'
         '• 公關文案化解與溝通對策：____________________________________________________\n'
         '  __________________________________________________________________________________')
    ]

    for idx, (label, content) in enumerate(feedback_tests, start=1):
        t4.cell(idx, 0).paragraphs[0].add_run(label).bold = True
        set_cell_format(t4.cell(idx, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")
        
        p = t4.cell(idx, 1).paragraphs[0]
        p.paragraph_format.line_spacing = 1.3
        p.add_run(content)
        set_cell_format(t4.cell(idx, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=100, pad_bottom=100)

    # ------------------ 六、 本週提示工程調校心得與反思（收斂總結） ------------------
    add_heading_paragraph('六、 本週提示工程調校心得與反思（收斂總結）')
    t5 = doc.add_table(rows=2, cols=2)
    set_table_props(t5)

    t5.cell(0, 0).paragraphs[0].add_run('提示工程調校\n重大突破\n\n【100字以內描述】').bold = True
    set_cell_format(t5.cell(0, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_rf1 = t5.cell(0, 1).paragraphs[0]
    p_rf1.paragraph_format.line_spacing = 1.3
    p_rf1.add_run('請具體說明：在調校 System Instructions 與 Few-shot 示範的過程中，你們發現了哪一項「最有效的調優關鍵技巧」？如何讓原本失焦或罐頭式的回覆變得精準可用？\n\n\n\n')
    set_cell_format(t5.cell(0, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=110, pad_bottom=110)

    t5.cell(1, 0).paragraphs[0].add_run('跨模態提示詞\n延伸構想\n\n【邁向第四週】').bold = True
    set_cell_format(t5.cell(1, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_rf2 = t5.cell(1, 1).paragraphs[0]
    p_rf2.paragraph_format.line_spacing = 1.3
    p_rf2.add_run('本週建立的 CRISP 角色與場景語境，預計在第四週如何轉化為 Google Flow（Nano Banana Pro / Imagen 4）的生圖提示詞（S-S-L-M 架構：主體、場景、光影、媒材）？\n\n\n\n')
    set_cell_format(t5.cell(1, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=110, pad_bottom=110)

    # ------------------ 七、 優質示範參考標竿（Reference Benchmark） ------------------
    add_heading_paragraph('七、 優質示範參考標竿（Reference Benchmark）', color=COLOR_BENCHMARK, size_pt=10.5, space_before=12, space_after=3)
    
    p_bm_title = doc.add_paragraph()
    p_bm_title.paragraph_format.space_after = Pt(2)
    r_bm_t = p_bm_title.add_run('【醜蔬果剩食拯救專案《格外美好》之 CRISP 與 Few-shot 示範】')
    r_bm_t.bold = True
    r_bm_t.font.size = Pt(9.0)
    r_bm_t.font.color.rgb = COLOR_BENCHMARK

    bm_bullets = [
        '• 核心定位：針對 18~24 歲外宿大學生，將產地格外品蔬果轉化為平價五星級深夜宵夜。',
        '• CRISP 調校關鍵：\n'
        '  - Context: 台灣每年數萬噸蔬果因外觀規格不符遭產地掩埋；大學生月底伙食費窘迫、缺乏蔬果纖維。\n'
        '  - Role: 親切又帶點毒舌幽默的深夜外宿料理主廚兼永續生活推手。\n'
        '  - Instruction: 撰寫 3 組切入不同外宿痛點的 IG 貼文文案，每組包含標題、內文與 CTA。\n'
        '  - Scenario: 月底看著錢包嘆氣、叫外送又覺得又貴又油膩的疲憊大學生。\n'
        '  - Product: 採用 Markdown 表格，精確輸出【次受眾類型】、【Hook 標題】、【80 字痛點內文】與【CTA】。',
        '• Few-shot 抗拒化解示範：\n'
        '  - 質疑：「醜蔬果是不是快爛掉的瑕疵品？吃了會不會拉肚子？」\n'
        '  - 意圖：食安品質疑慮 ｜ 情緒：強烈負面 ｜ 策略：拍攝農產專家盲測糖度影片，主打「醜在表皮、甜在核心、通過產銷履歷農藥殘留零檢出」。'
    ]

    for b in bm_bullets:
        p_b = doc.add_paragraph()
        p_b.paragraph_format.space_after = Pt(2)
        p_b.paragraph_format.line_spacing = 1.2
        r_b = p_b.add_run(b)
        r_b.font.size = Pt(8.5)
        r_b.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    # ------------------ 八、 課堂即時諮詢與審查回饋欄（教師 / 助教填寫） ------------------
    add_heading_paragraph('八、 課堂即時諮詢與審查回饋欄（教師 / 助教填寫）', color=COLOR_REVIEW, size_pt=10.5, space_before=12, space_after=3)
    t6 = doc.add_table(rows=2, cols=2)
    set_table_props(t6)

    t6.cell(0, 0).paragraphs[0].add_run('提示工程深度\n評核狀態').bold = True
    set_cell_format(t6.cell(0, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_st = t6.cell(0, 1).paragraphs[0]
    p_st.paragraph_format.line_spacing = 1.3
    p_st.add_run('☐ 審核通過：CRISP 結構嚴謹，情境生動具象，Few-shot 分類穩定且對策深刻。\n')
    p_st.add_run('☐ 建議微調：受眾痛點略顯泛化，或分類任務溫度偏高，建議修正限制條件。\n')
    p_st.add_run('☐ 重新調校：缺乏具體場景（Scenario）或輸出規格抽象，請重構提示詞後再測。')
    set_cell_format(t6.cell(0, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=100, pad_bottom=100)

    t6.cell(1, 0).paragraphs[0].add_run('教師指導建議\n與簽章').bold = True
    set_cell_format(t6.cell(1, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_sig = t6.cell(1, 1).paragraphs[0]
    p_sig.paragraph_format.line_spacing = 1.3
    p_sig.add_run('指導意見：\n\n\n\n')
    p_sig.add_run('審查簽名：____________________    日期：2026 年 ______ 月 ______ 日')
    set_cell_format(t6.cell(1, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=100, pad_bottom=100)

    # 儲存檔案
    output_path = 'SDGs_結構化提示工程與多情境資料擴展實作單.docx'
    doc.save(output_path)
    print(f'Successfully generated {output_path}')

if __name__ == '__main__':
    create_document()
