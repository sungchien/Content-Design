import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = docx.Document()
    
    # 頁面邊界設定 (2.032 cm / 0.8 inch 全等邊界，維持現有 Word 實作單規格)
    section = doc.sections[0]
    section.top_margin = Cm(2.032)
    section.bottom_margin = Cm(2.032)
    section.left_margin = Cm(2.032)
    section.right_margin = Cm(2.032)
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    
    # 全域樣式設定
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Microsoft JhengHei'
    normal_style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    normal_style.font.size = Pt(9.5)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(3)

    # 樣式常數定義
    COLOR_PRIMARY = RGBColor(0x1a, 0x56, 0xb8)     # 皇家藍 (標題與章節標題)
    COLOR_MUTED = RGBColor(0x7f, 0x8c, 0x8d)       # 灰字 (頂部引導備註)
    COLOR_BODY = RGBColor(0x55, 0x55, 0x55)        # 說明次要字
    COLOR_BENCHMARK = RGBColor(0x2e, 0x7d, 0x32)   # 墨綠色 (示範標竿)
    COLOR_REVIEW = RGBColor(0xc0, 0x39, 0x2b)      # 緋紅色 (教師審查欄)
    
    FILL_HEADER = "F0F4F8"                         # 欄位標籤淺藍灰底色
    BORDER_COLOR = "D0D7DE"                        # 表格細邊框淺灰色
    
    TOTAL_WIDTH_DXA = 9936                         # 17.526 cm (可列印區域總寬度)
    LABEL_WIDTH_DXA = 2880                         # 5.08 cm
    CONTENT_WIDTH_DXA = 7056                       # 12.446 cm
    
    def set_cell_format(cell, width_dxa=None, fill_hex=None, v_align="center", pad_top=100, pad_bottom=100, pad_left=130, pad_right=130):
        tcPr = cell._tc.get_or_add_tcPr()
        
        # 寬度
        if width_dxa is not None:
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW')
                tcPr.append(tcW)
            tcW.set(qn('w:type'), 'dxa')
            tcW.set(qn('w:w'), str(width_dxa))
            
        # 垂直對齊
        vAlign = tcPr.find(qn('w:vAlign'))
        if vAlign is None:
            vAlign = OxmlElement('w:vAlign')
            tcPr.append(vAlign)
        vAlign.set(qn('w:val'), v_align)
        
        # 內距 padding
        tcMar = tcPr.find(qn('w:tcMar'))
        if tcMar is None:
            tcMar = OxmlElement('w:tcMar')
            tcPr.append(tcMar)
        for side, pad in [('top', pad_top), ('bottom', pad_bottom), ('left', pad_left), ('right', pad_right)]:
            m = tcMar.find(qn(f'w:{side}'))
            if m is None:
                m = OxmlElement(f'w:{side}')
                tcMar.append(m)
            m.set(qn('w:w'), str(pad))
            m.set(qn('w:type'), 'dxa')
            
        # 邊框
        borders_xml = f'''
        <w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>
        </w:tcBorders>
        '''
        existing_borders = tcPr.find(qn('w:tcBorders'))
        if existing_borders is not None:
            tcPr.remove(existing_borders)
        tcPr.append(parse_xml(borders_xml))
        
        # 底色
        if fill_hex:
            shd = tcPr.find(qn('w:shd'))
            if shd is None:
                shd = OxmlElement('w:shd')
                tcPr.append(shd)
            shd.set(qn('w:fill'), fill_hex)

    def set_table_props(table):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = table._tbl.tblPr
        tblW = tblPr.find(qn('w:tblW'))
        if tblW is not None:
            tblW.set(qn('w:type'), 'auto')
            tblW.set(qn('w:w'), '0')
        # 防止跨頁截斷行
        for row in table.rows:
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml('<w:cantSplit xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))

    def add_heading_paragraph(text, color=COLOR_PRIMARY, size_pt=11.5, space_before=10, space_after=3):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Microsoft JhengHei'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
        run.font.size = Pt(size_pt)
        run.font.bold = True
        run.font.color.rgb = color
        return p

    # ------------------ 文件頂部抬頭 ------------------
    p0 = doc.add_paragraph()
    p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p0.paragraph_format.space_after = Pt(2)
    r0 = p0.add_run('數位內容設計（Digital Content Design）｜ PBL 專題實作導引（第四週）')
    r0.font.name = 'Microsoft JhengHei'
    r0._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    r0.font.size = Pt(9.0)
    r0.font.color.rgb = COLOR_MUTED

    p1 = doc.add_paragraph()
    p1.paragraph_format.space_after = Pt(4)
    p1.paragraph_format.space_before = Pt(0)
    r1 = p1.add_run('第四週：SDGs 專題核心主視覺生成與模型調校實作單')
    r1.font.name = 'Microsoft JhengHei'
    r1._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    r1.font.size = Pt(17.0)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_PRIMARY

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(8)
    p2.paragraph_format.line_spacing = 1.25
    r2 = p2.add_run(
        '本實作單用於第四週課堂工作坊（60 分鐘）。請各組延續第三週所淬鍊之 SDGs 核心倡議主題，實踐「雙階段多模態視覺產製工作流」：'
        '首先於 Google AI Studio（aistudio.google.com）透過 Gemini 3.1 Pro 與 CRISP 架構產出 S-S-L-M 提示詞，'
        '並以 Nano Banana 2 Lite 進行低延遲、秒級概念驗證與 A/B 提案測試；接著將篩選出之滿意草圖與提示詞轉移至 Google Flow，'
        '運用 Nano Banana 2 的「以圖生圖（Image-to-Image）」與「筆刷局部修復（Interactive Inpainting）」進行高畫質升級與微觀精修。'
        '完成後於課堂結束前提交，經教師/助教諮詢查核後，作為專案期末倡議海報與社群主視覺之核心資產。'
    )
    r2.font.name = 'Microsoft JhengHei'
    r2._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = COLOR_BODY

    # ------------------ 一、 小組基本資訊與多模態視覺職能分工 ------------------
    add_heading_paragraph('一、 小組基本資訊與多模態視覺職能分工')
    t0 = doc.add_table(rows=3, cols=4)
    set_table_props(t0)
    col_w_t0 = 2484  # 9936 / 4 = 2484 dxa
    
    # R0
    t0.cell(0, 0).paragraphs[0].add_run('專案組別').bold = True
    t0.cell(0, 1).paragraphs[0].add_run('第 _______ 組')
    t0.cell(0, 2).paragraphs[0].add_run('小組組名').bold = True
    t0.cell(0, 3).paragraphs[0].add_run('《 ___________________ 》')
    set_cell_format(t0.cell(0, 0), width_dxa=col_w_t0, fill_hex=FILL_HEADER)
    set_cell_format(t0.cell(0, 1), width_dxa=col_w_t0)
    set_cell_format(t0.cell(0, 2), width_dxa=col_w_t0, fill_hex=FILL_HEADER)
    set_cell_format(t0.cell(0, 3), width_dxa=col_w_t0)

    # R1
    t0.cell(1, 0).paragraphs[0].add_run('專案暫定名稱').bold = True
    t0.cell(1, 1).paragraphs[0].add_run('《 _______________________________ 》')
    t0.cell(1, 2).paragraphs[0].add_run('選定 SDGs 目標').bold = True
    t0.cell(1, 3).paragraphs[0].add_run('第 【    】 號目標：_________________')
    set_cell_format(t0.cell(1, 0), width_dxa=col_w_t0, fill_hex=FILL_HEADER)
    set_cell_format(t0.cell(1, 1), width_dxa=col_w_t0)
    set_cell_format(t0.cell(1, 2), width_dxa=col_w_t0, fill_hex=FILL_HEADER)
    set_cell_format(t0.cell(1, 3), width_dxa=col_w_t0)

    # R2 (合併後 3 格)
    t0.cell(2, 0).paragraphs[0].add_run('今日多模態生成\n職能分工').bold = True
    set_cell_format(t0.cell(2, 0), width_dxa=col_w_t0, fill_hex=FILL_HEADER)
    
    c_merged = t0.cell(2, 1).merge(t0.cell(2, 3))
    p_roles = c_merged.paragraphs[0]
    p_roles.paragraph_format.line_spacing = 1.3
    p_roles.add_run('• 提示工程師 (Prompt Engineer / AI Studio 操控)：____________________（負責 CRISP 視覺腳本建構、S-S-L-M 四維度拆解與 AI Studio 參數設定）\n')
    p_roles.add_run('• 藝術總監 / 設計師 (Art Director / Google Flow 操控)：________________（負責構圖光影審查、Flow 參考圖錨定強度 0.45、筆刷遮罩局部修復 Inpainting）\n')
    p_roles.add_run('• 永續內容策略員 (Content Strategist / 倡議把關)：____________________（負責核對畫面痛點共感度、生活化微觀視角、自然語言排除瑕疵條件把關）')
    set_cell_format(c_merged, width_dxa=col_w_t0 * 3, v_align="top", pad_top=120, pad_bottom=120)

    # ------------------ 二、 階段一：Google AI Studio 視覺敘事腳本與 S-S-L-M 提示生成 ------------------
    add_heading_paragraph('二、 階段一：Google AI Studio 視覺敘事腳本與 S-S-L-M 提示生成')
    t1 = doc.add_table(rows=3, cols=2)
    set_table_props(t1)

    # R0: AI Studio 參數配置
    t1.cell(0, 0).paragraphs[0].add_run('AI Studio 介面參數\n與全域提示詞設定\n\n【靈魂與發散度】').bold = True
    set_cell_format(t1.cell(0, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_ai_prm = t1.cell(0, 1).paragraphs[0]
    p_ai_prm.paragraph_format.line_spacing = 1.3
    p_ai_prm.add_run('1. System Instructions 全域角色設定（專家靈魂）：\n').font.color.rgb = COLOR_BODY
    p_ai_prm.add_run('   你是專精於【 ______________________________________ 】的資深藝術總監兼提示架構師。\n')
    p_ai_prm.add_run('   你的職責是將永續議題轉化為具備強烈情感張力的影像提示詞，嚴格依循 S-S-L-M 四維度，\n')
    p_ai_prm.add_run('   並以中文自然語言輸出，將畫質排斥特徵無縫融入提示詞尾端。\n')
    p_ai_prm.add_run('2. 運算參數調校：Temperature 設定值：【 0.7 】 ｜ Top-P 設定值：【 0.95 】\n')
    p_ai_prm.add_run('   • 參數設計理由：維持視覺描述的文學修辭與生動想像力，同時兼顧提示詞語法規範與收斂度。')
    set_cell_format(t1.cell(0, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=110, pad_bottom=110)

    # R1: CRISP 視覺腳本提示詞
    t1.cell(1, 0).paragraphs[0].add_run('CRISP 視覺腳本\n提示詞結構設計\n\n【產出 3 組候選提示】').bold = True
    set_cell_format(t1.cell(1, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_crisp = t1.cell(1, 1).paragraphs[0]
    p_crisp.paragraph_format.line_spacing = 1.3
    p_crisp.add_run('小組輸入至 Google AI Studio 的 CRISP 提示詞內容：\n').font.color.rgb = COLOR_BODY
    p_crisp.add_run('• Context (語境)：_______________________________________________________________\n')
    p_crisp.add_run('• Role (角色)：獲獎無數的國家地理雜誌資深攝影總監兼永續視覺敘事專家。\n')
    p_crisp.add_run('• Instruction (核心任務)：運用 S-S-L-M 架構，針對本議題產出 3 組切入不同視角、可用於 Nano Banana 2 Lite 的繁體中文生圖提示詞。\n')
    p_crisp.add_run('• Scenario (情境痛點)：___________________________________________________________\n')
    p_crisp.add_run('• Product (規格限制)：每組提示詞約 120~180 字，包含完整 S-S-L-M 維度，並於尾端以自然語言融入畫質排斥條件（杜絕塑膠感、多餘肢體、失真瑕疵）。')
    set_cell_format(t1.cell(1, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=110, pad_bottom=110)

    # R2: S-S-L-M 四維度拆解分析
    t1.cell(2, 0).paragraphs[0].add_run('S-S-L-M\n四維度結構拆解\n\n【核心控制面板】').bold = True
    set_cell_format(t1.cell(2, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_sslm = t1.cell(2, 1).paragraphs[0]
    p_sslm.paragraph_format.line_spacing = 1.3
    p_sslm.add_run('請將小組擬定之核心視覺畫面拆解為四層結構控制語：\n').font.color.rgb = COLOR_BODY
    p_sslm.add_run('• S (Subject 主體與微表情)：_____________________________________________________\n')
    p_sslm.add_run('• S (Setting & Composition 場景與構圖)：_________________________________________\n')
    p_sslm.add_run('• L (Lighting & Mood 光影氛圍與色彩)：___________________________________________\n')
    p_sslm.add_run('• M (Medium & Technical Specs 媒材與光學規格)：_________________________________')
    set_cell_format(t1.cell(2, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=110, pad_bottom=110)

    # ------------------ 三、 階段二：Nano Banana 2 Lite 快速概念驗證（AI Studio 實測） ------------------
    add_heading_paragraph('三、 階段二：Nano Banana 2 Lite 快速概念驗證（AI Studio 實測）')
    t2 = doc.add_table(rows=4, cols=2)
    set_table_props(t2)

    # R0: 3組切角提案
    t2.cell(0, 0).paragraphs[0].add_run('3 組切角提案\n生成提示詞與概念\n\n【快速原型 A/B 測試】').bold = True
    set_cell_format(t2.cell(0, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_props = t2.cell(0, 1).paragraphs[0]
    p_props.paragraph_format.line_spacing = 1.3
    p_props.add_run('【提案一（微觀受害特寫 / 情感共鳴）】：\n').bold = True
    p_props.add_run('提示詞：_________________________________________________________________________\n\n')
    p_props.add_run('【提案二（宏觀環境浩劫 / 生態震撼）】：\n').bold = True
    p_props.add_run('提示詞：_________________________________________________________________________\n\n')
    p_props.add_run('【提案三（象徵意象與光影對比 / 警示啟發）】：\n').bold = True
    p_props.add_run('提示詞：_________________________________________________________________________')
    set_cell_format(t2.cell(0, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=100, pad_bottom=100)

    # R1: 主測提案選定
    t2.cell(1, 0).paragraphs[0].add_run('小組主測提案選定\n與規格參數').bold = True
    set_cell_format(t2.cell(1, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_sel = t2.cell(1, 1).paragraphs[0]
    p_sel.paragraph_format.line_spacing = 1.3
    p_sel.add_run('• 團隊選定之最優提案：☐ 提案一    ☐ 提案二    ☐ 提案三\n')
    p_sel.add_run('• 選定理由（為何此畫面最能引發目標受眾痛點共鳴？約 40 字）：\n  __________________________________________________________________________________\n')
    p_sel.add_run('• 生成參數配置：長寬比 Output Aspect：☐ 4:5（直式社群/海報）  ☐ 16:9（橫式簡報）  ☐ 1:1\n')
    p_sel.add_run('• 自然語言排除特徵：畫面杜絕 3D 動畫卡通感、塑膠平滑皮膚、多餘畸形肢體、低解析度模糊瑕疵。')
    set_cell_format(t2.cell(1, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=100, pad_bottom=100)

    # R2: 草圖黏貼與診斷
    t2.cell(2, 0).paragraphs[0].add_run('Nano Banana 2 Lite\n概念草圖黏貼處\n\n【秒級原型成果】').bold = True
    set_cell_format(t2.cell(2, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_img1 = t2.cell(2, 1).paragraphs[0]
    p_img1.paragraph_format.line_spacing = 1.3
    p_img1.add_run('【請於此處貼上由 Nano Banana 2 Lite 生成之滿意概念草圖截圖（建議寬度約 8~10 cm）】\n\n\n\n\n\n\n\n')
    p_img1.add_run('圖檔命名：Team_XX_Lite_Prototype.png ｜ 生成耗時：約 ______ 秒')
    set_cell_format(t2.cell(2, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=100, pad_bottom=100)

    # R3: 原型優點與缺陷診斷
    t2.cell(3, 0).paragraphs[0].add_run('原型優點與\n待改善缺陷診斷').bold = True
    set_cell_format(t2.cell(3, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_diag = t2.cell(3, 1).paragraphs[0]
    p_diag.paragraph_format.line_spacing = 1.3
    p_diag.add_run('• 值得保留之優點（如主體姿態、光影構圖）：_______________________________________\n')
    p_diag.add_run('• 亟待改善之缺陷（如材質細節、邊緣瑕疵、特定局部扭曲）：\n  __________________________________________________________________________________')
    set_cell_format(t2.cell(3, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=100, pad_bottom=100)

    # ------------------ 四、 階段三：Google Flow 高階模型調校與精修（Nano Banana 2 實戰） ------------------
    add_heading_paragraph('四、 階段三：Google Flow 高階模型調校與精修（Nano Banana 2 實戰）')
    t3 = doc.add_table(rows=4, cols=2)
    set_table_props(t3)

    # R0: Google Flow 設定
    t3.cell(0, 0).paragraphs[0].add_run('Google Flow 模型\n與參考圖錨定設定\n\n【以圖生圖配置】').bold = True
    set_cell_format(t3.cell(0, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_flow_cfg = t3.cell(0, 1).paragraphs[0]
    p_flow_cfg.paragraph_format.line_spacing = 1.3
    p_flow_cfg.add_run('1. 選用模型版本：【 Nano Banana 2 】（Google Flow 專業創作者工作流）\n')
    p_flow_cfg.add_run('2. 參考圖片載入 (Image Reference)：已載入階段二 Nano Banana 2 Lite 滿意草圖\n')
    p_flow_cfg.add_run('3. 重繪 / 變化強度 (Variation Strength)：設定值【 0.45 】（建議值 0.4 ~ 0.5）\n')
    p_flow_cfg.add_run('   • 參數設計依據：嚴格繼承原草圖之透視幾何與主體受困姿態，賦予模型重算微觀質感的運算空間。')
    set_cell_format(t3.cell(0, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=100, pad_bottom=100)

    # R1: 筆刷遮罩局部修復 (Inpainting)
    t3.cell(1, 0).paragraphs[0].add_run('筆刷遮罩局部修復\n(Interactive Inpainting)\n\n【精確無縫置換】').bold = True
    set_cell_format(t3.cell(1, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_inp = t3.cell(1, 1).paragraphs[0]
    p_inp.paragraph_format.line_spacing = 1.3
    p_inp.add_run('• 啟用工具：Google Flow 畫布筆刷工具（Brush / Mask Magic）\n')
    p_inp.add_run('• 塗抹修復之特定區域（例如：受困漁網粗細、背景水質雜質、肢體鱗片紋理）：\n  __________________________________________________________________________________\n')
    p_inp.add_run('• 局部修復提示詞（Inpainting Prompt，精確指定替換特徵）：\n  __________________________________________________________________________________')
    set_cell_format(t3.cell(1, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=100, pad_bottom=100)

    # R2: Flow 整合提示詞
    t3.cell(2, 0).paragraphs[0].add_run('Google Flow 最終\n結構化提示詞\n\n【自然語言排除特徵】').bold = True
    set_cell_format(t3.cell(2, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_fl_prm = t3.cell(2, 1).paragraphs[0]
    p_fl_prm.paragraph_format.line_spacing = 1.3
    p_fl_prm.add_run('（請填寫輸入至 Google Flow 之完整提示詞，無獨立負向輸入框，需將排斥特徵融入末尾）：\n').font.color.rgb = COLOR_BODY
    p_fl_prm.add_run('_____________________________________________________________________________________\n')
    p_fl_prm.add_run('_____________________________________________________________________________________\n')
    p_fl_prm.add_run('_____________________________________________________________________________________\n')
    p_fl_prm.add_run('排斥條件句：畫面杜絕 _______________________________________________________________')
    set_cell_format(t3.cell(2, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=100, pad_bottom=100)

    # R3: 最終成圖黏貼
    t3.cell(3, 0).paragraphs[0].add_run('最終精修主視覺成圖\n(Nano Banana 2)\n\n【高解析倡議成圖】').bold = True
    set_cell_format(t3.cell(3, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_img2 = t3.cell(3, 1).paragraphs[0]
    p_img2.paragraph_format.line_spacing = 1.3
    p_img2.add_run('【請於此處貼上經 Google Flow 精修與 Inpainting 後之最終主視覺成圖（建議寬度約 8~10 cm）】\n\n\n\n\n\n\n\n')
    p_img2.add_run('圖檔命名：Team_XX_Final_KeyVisual.png ｜ 影像規格：4:5 直式高畫質')
    set_cell_format(t3.cell(3, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=100, pad_bottom=100)

    # ------------------ 五、 迭代調校過程反思與排錯記錄 ------------------
    add_heading_paragraph('五、 迭代調校過程反思與排錯記錄')
    t4 = doc.add_table(rows=3, cols=2)
    set_table_props(t4)

    # R0: 第一次迭代
    t4.cell(0, 0).paragraphs[0].add_run('第一次迭代反思\n【草圖階段調校】').bold = True
    set_cell_format(t4.cell(0, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_it1 = t4.cell(0, 1).paragraphs[0]
    p_it1.paragraph_format.line_spacing = 1.3
    p_it1.add_run('• 發現之核心缺陷（如構圖失焦、語意污染、畫面過於卡通）：\n  __________________________________________________________________________________\n')
    p_it1.add_run('• 提示詞/參數具體修改動作：_____________________________________________________\n')
    p_it1.add_run('• 改善成果評價：_________________________________________________________________')
    set_cell_format(t4.cell(0, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=100, pad_bottom=100)

    # R1: 第二次迭代
    t4.cell(1, 0).paragraphs[0].add_run('第二次迭代反思\n【精修階段調校】').bold = True
    set_cell_format(t4.cell(1, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_it2 = t4.cell(1, 1).paragraphs[0]
    p_it2.paragraph_format.line_spacing = 1.3
    p_it2.add_run('• 發現之核心缺陷（如參考圖權重過大無法修正、局部修復邊緣不自然）：\n  __________________________________________________________________________________\n')
    p_it2.add_run('• 提示詞/參數/筆刷遮罩修改動作：_________________________________________________\n')
    p_it2.add_run('• 改善成果評價：_________________________________________________________________')
    set_cell_format(t4.cell(1, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=100, pad_bottom=100)

    # R2: 跨工具協同體會
    t4.cell(2, 0).paragraphs[0].add_run('雙階段多模態工作流\n整體收斂體會').bold = True
    set_cell_format(t4.cell(2, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_it3 = t4.cell(2, 1).paragraphs[0]
    p_it3.paragraph_format.line_spacing = 1.3
    p_it3.add_run('對比「直接在單一工具抽卡盲測」，你們認為「先在 AI Studio 用 Nano Banana 2 Lite 快速低成本驗證靈感，再轉入 Google Flow 用 Nano Banana 2 精修與局部替換」的工作流帶來了什麼關鍵優勢？（約 60 字）：\n\n\n')
    set_cell_format(t4.cell(2, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=100, pad_bottom=100)

    # ------------------ 六、 本週核心觀念自我檢核清單（Checklist） ------------------
    add_heading_paragraph('六、 本週核心觀念自我檢核清單（Checklist）')
    t5 = doc.add_table(rows=1, cols=2)
    set_table_props(t5)

    t5.cell(0, 0).paragraphs[0].add_run('多模態生成能力\n自我指標檢核').bold = True
    set_cell_format(t5.cell(0, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_ck = t5.cell(0, 1).paragraphs[0]
    p_ck.paragraph_format.line_spacing = 1.35
    p_ck.add_run('完成本週實作後，請組員共同勾選確認已掌握之專業技能：\n').font.color.rgb = COLOR_BODY
    p_ck.add_run('☐ 擴散模型原理：能以自己的話向組員解釋「正向加噪」與「反向去噪雕琢」的生成邏輯。\n')
    p_ck.add_run('☐ S-S-L-M 架構掌握：能自覺拆解出主體（S）、場景構圖（S）、光影色彩（L）、媒材規格（M）四維度。\n')
    p_ck.add_run('☐ 工具分工認知：清楚說明 Nano Banana 2 Lite（秒級原型驗證）與 Google Flow Nano Banana 2（高精以圖生圖與局部重繪）的定位差異。\n')
    p_ck.add_run('☐ 自然語言負向控制：能在無獨立負向提示詞框之介面中，以自然語言流暢融入排除條件。\n')
    p_ck.add_run('☐ 局部修復實務：能運用 Google Flow 筆刷遮罩（Inpainting）在不動及主體的前提下修復微觀細節瑕疵。')
    set_cell_format(t5.cell(0, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=110, pad_bottom=110)

    # ------------------ 七、 優質示範參考標竿（Reference Benchmark） ------------------
    add_heading_paragraph('七、 優質示範參考標竿（Reference Benchmark）', color=COLOR_BENCHMARK, size_pt=10.5, space_before=12, space_after=3)
    
    p_bm_title = doc.add_paragraph()
    p_bm_title.paragraph_format.space_after = Pt(2)
    r_bm_t = p_bm_title.add_run('【海洋廢棄物浩劫《幼龜受困海廢》之雙階段多模態生成示範全紀錄】')
    r_bm_t.bold = True
    r_bm_t.font.size = Pt(9.0)
    r_bm_t.font.color.rgb = COLOR_BENCHMARK

    bm_bullets = [
        '• 核心倡議主題：SDG 14 保育海洋生態——以「幽靈漁網纏繞幼龜」傳達海洋塑膠廢棄物對微觀生命體的致命浩劫。',
        '• 階段一 Google AI Studio (Gemini 3.1 Pro 視覺腳本) 調校：\n'
        '  - System Instructions: 賦予國家地理雜誌資深攝影總監專家靈魂，Temperature 0.7，Top-P 0.95。\n'
        '  - 輸出產物：成功產出 3 組切入不同視角（幼龜特寫、珊瑚礁浩劫、微塑膠風暴）之中文生圖提示詞。',
        '• 階段二 Google AI Studio (Nano Banana 2 Lite 快速原型驗證)：\n'
        '  - 實測「提案一：幼龜受困特寫」，Output Aspect 設定 4:5，次秒級產出低延遲概念草圖，確立海龜仰角眼神與水下透視之動人張力。',
        '• 階段三 Google Flow (Nano Banana 2 高精精修與 Inpainting)：\n'
        '  - 以圖生圖錨定：載入 Lite 原型草圖，重繪強度設為 0.45，完美保全原海龜神態與空間幾何。\n'
        '  - Inpainting 局部修復：以筆刷塗抹漁網線條，提示詞指定「增強水下殘破尼龍網線的風化纖維質感與微流水泡」，實現高解析微觀置換。\n'
        '  - 最終提示詞自然語言排除：尾端加入「畫面杜絕 3D 動畫卡通感、塑膠平滑皮膚、多餘鰭肢或雙頭畸形，無人造乾淨泳池水質感，無低解析度模糊瑕疵」。'
    ]

    for b in bm_bullets:
        p_b = doc.add_paragraph()
        p_b.paragraph_format.space_after = Pt(2)
        p_b.paragraph_format.line_spacing = 1.2
        r_b = p_b.add_run(b)
        r_b.font.size = Pt(8.5)
        r_b.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    # ------------------ 八、 課堂即時諮詢與審查回饋欄（教師 / 助教填寫） ------------------
    add_heading_paragraph('八、 課堂即時諮詢與審查回饋欄（教師 / 助教填寫）', color=COLOR_REVIEW, size_pt=10.5, space_before=12, space_after=3)
    t6 = doc.add_table(rows=2, cols=2)
    set_table_props(t6)

    t6.cell(0, 0).paragraphs[0].add_run('多模態視覺深度\n評核狀態').bold = True
    set_cell_format(t6.cell(0, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_st = t6.cell(0, 1).paragraphs[0]
    p_st.paragraph_format.line_spacing = 1.3
    p_st.add_run('☐ 審核通過：S-S-L-M 架構嚴謹，草圖與精修圖對照鮮明，Inpainting 局部修復自然，倡議共感強烈。\n')
    p_st.add_run('☐ 建議微調：光影氛圍或材質描述略顯抽象，或以圖生圖重繪強度偏離建議值，建議微調後再測。\n')
    p_st.add_run('☐ 重新調校：畫面出現明顯失真/多餘肢體，或未落實自然語言排除條件，請重修提示詞。')
    set_cell_format(t6.cell(0, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=100, pad_bottom=100)

    t6.cell(1, 0).paragraphs[0].add_run('教師指導建議\n與簽章').bold = True
    set_cell_format(t6.cell(1, 0), width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER, v_align="center")

    p_sig = t6.cell(1, 1).paragraphs[0]
    p_sig.paragraph_format.line_spacing = 1.3
    p_sig.add_run('指導意見：\n\n\n\n')
    p_sig.add_run('審查簽名：____________________    日期：2026 年 ______ 月 ______ 日')
    set_cell_format(t6.cell(1, 1), width_dxa=CONTENT_WIDTH_DXA, v_align="top", pad_top=100, pad_bottom=100)

    # 儲存檔案
    output_path = 'SDGs_專題核心主視覺生成與模型調校實作單.docx'
    doc.save(output_path)
    print(f'Successfully generated {output_path}')

if __name__ == '__main__':
    create_document()
