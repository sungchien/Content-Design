import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = docx.Document()
    
    # 頁面邊界設定 (2.032 cm / 0.8 inch 全等邊界)
    section = doc.sections[0]
    section.top_margin = Cm(2.032)
    section.bottom_margin = Cm(2.032)
    section.left_margin = Cm(2.032)
    section.right_margin = Cm(2.032)
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    
    # 全域樣式設定
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Microsoft JhengHei'
    normal_style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    normal_style.font.size = Pt(9.5)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(3)

    # 樣式常數定義
    COLOR_PRIMARY = RGBColor(0x1a, 0x56, 0xb8)     # 皇家藍 (標題與章節標題)
    COLOR_MUTED = RGBColor(0x7f, 0x8c, 0x8d)       # 灰字 (頂部引導備註)
    COLOR_BODY = RGBColor(0x55, 0x55, 0x55)        # 說明次要字
    COLOR_REVIEW = RGBColor(0xc0, 0x39, 0x2b)      # 緋紅色 (教師審查欄)
    
    FILL_HEADER = "F0F4F8"                         # 欄位標籤淺藍灰底色
    BORDER_COLOR = "D0D7DE"                        # 表格細邊框淺灰色
    
    LABEL_WIDTH_DXA = 2880                         # 5.08 cm
    CONTENT_WIDTH_DXA = 7056                       # 12.446 cm
    
    def set_cell_format(cell, width_dxa=None, fill_hex=None, v_align="center", pad_top=100, pad_bottom=100, pad_left=130, pad_right=130):
        tcPr = cell._tc.get_or_add_tcPr()
        
        # 寬度
        if width_dxa is not None:
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW')
                tcPr.append(tcW)
            tcW.set(qn('w:type'), 'dxa')
            tcW.set(qn('w:w'), str(width_dxa))
            
        # 垂直對齊
        vAlign = tcPr.find(qn('w:vAlign'))
        if vAlign is None:
            vAlign = OxmlElement('w:vAlign')
            tcPr.append(vAlign)
        vAlign.set(qn('w:val'), v_align)
        
        # 內距 padding
        tcMar = tcPr.find(qn('w:tcMar'))
        if tcMar is None:
            tcMar = OxmlElement('w:tcMar')
            tcPr.append(tcMar)
        for side, pad in [('top', pad_top), ('bottom', pad_bottom), ('left', pad_left), ('right', pad_right)]:
            m = tcMar.find(qn(f'w:{side}'))
            if m is None:
                m = OxmlElement(f'w:{side}')
                tcMar.append(m)
            m.set(qn('w:w'), str(pad))
            m.set(qn('w:type'), 'dxa')
            
        # 邊框
        borders_xml = f'''
        <w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>
        </w:tcBorders>
        '''
        existing_borders = tcPr.find(qn('w:tcBorders'))
        if existing_borders is not None:
            tcPr.remove(existing_borders)
        tcPr.append(parse_xml(borders_xml))
        
        # 底色
        if fill_hex:
            shd = tcPr.find(qn('w:shd'))
            if shd is None:
                shd = OxmlElement('w:shd')
                tcPr.append(shd)
            shd.set(qn('w:fill'), fill_hex)

    def set_table_props(table):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = table._tbl.tblPr
        tblW = tblPr.find(qn('w:tblW'))
        if tblW is not None:
            tblW.set(qn('w:type'), 'auto')
            tblW.set(qn('w:w'), '0')
        for row in table.rows:
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml('<w:cantSplit xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))

    def add_heading_paragraph(text, color=COLOR_PRIMARY, size_pt=11.5, space_before=10, space_after=3):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Microsoft JhengHei'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
        run.font.size = Pt(size_pt)
        run.font.bold = True
        run.font.color.rgb = color
        return p

    # ------------------ 文件頂部抬頭 ------------------
    p0 = doc.add_paragraph()
    p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p0.paragraph_format.space_after = Pt(2)
    r0 = p0.add_run('數位內容設計（Digital Content Design）｜ PBL 專題實作導引（第七週）')
    r0.font.name = 'Microsoft JhengHei'
    r0._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    r0.font.size = Pt(9.0)
    r0.font.color.rgb = COLOR_MUTED

    p1 = doc.add_paragraph()
    p1.paragraph_format.space_after = Pt(4)
    p1.paragraph_format.space_before = Pt(0)
    r1 = p1.add_run('第七週：SDGs 專題生成內容倫理偏誤審計與 SynthID 防偽查核單')
    r1.font.name = 'Microsoft JhengHei'
    r1._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    r1.font.size = Pt(17.0)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_PRIMARY

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(8)
    p2.paragraph_format.line_spacing = 1.25
    r2 = p2.add_run(
        '本實作單用於第七週課堂工作坊（60 分鐘）。本週作為期中考前「概念驗證階段」之收官拼圖，請各組結合第四週主視覺海報、第五週 Campaign 文案與第六週智財合規成果，全面執行生成內容之社會倫理與真實性查核：'
        '首先透過「公立小學 6 歲女孩的 25 歲 AI 未來照」案例進行演算法偏見與想像力殖民之深度思辨，並執行預設提示詞偏誤誘發實驗；'
        '接著排查專案是否誤觸「貧窮色情（Poverty Porn）」、「救世主情結」與「視覺漂綠（Greenwashing）」三大倫理地雷，運用 S-S-L-M 去偏誤與平視賦權提示工程完成重構；'
        '隨後透過 Google SynthID 與反向圖片搜尋（About this image）檢驗自有與外部素材之防偽溯源，擬定正式《負責任 AI 創作與透明揭露聲明》；'
        '最後完成課堂 3 分鐘倫理紅隊同儕互評，經教師／助教查核簽章後，作為期中與期末專題之正式倫理合規依據。'
    )
    r2.font.name = 'Microsoft JhengHei'
    r2._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    r2.font.size = Pt(9.0)
    r2.font.color.rgb = COLOR_BODY

    # ------------------ 表格一：基本資訊與既有資產 ------------------
    add_heading_paragraph('【模組一：專案基本資訊與第 4～6 週既有資產回顧】')
    
    t1 = doc.add_table(rows=4, cols=2)
    set_table_props(t1)
    
    fields_t1 = [
        ('小組組別與專案名稱', '組別：第 ______ 組 ｜ 專案名稱：《__________________________________________________》'),
        ('選定之 SDGs 目標', '第 ______ 號核心目標：________________________（中文完整名稱）'),
        ('第 4～5 週既有資產回顧\n（主視覺人物與倡議標語）', '【第四週主視覺海報之人物／場景設定】：\n__________________________________________________________________________________\n【第五週 Campaign 一句話 Slogan】：\n《________________________________________________________________》'),
        ('小組成員與倫理審查分工', '倫理紅隊主審（PM）：______________ ｜ 視覺去偏誤重構（視覺）：______________\n反漂綠與文案排查（文案）：__________ ｜ SynthID 與反向搜尋查核：______________')
    ]
    
    for i, (label, val) in enumerate(fields_t1):
        c0 = t1.cell(i, 0)
        c1 = t1.cell(i, 1)
        set_cell_format(c0, width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER)
        set_cell_format(c1, width_dxa=CONTENT_WIDTH_DXA)
        
        p_lbl = c0.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(9.5)
        
        p_val = c1.paragraphs[0]
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.5)

    # ------------------ 表格二：案例思辨與演算法偏誤誘發實驗 ------------------
    add_heading_paragraph('【模組二：真實案例思辨（「6 歲女孩 AI 未來照」）與演算法偏誤誘發實驗】')
    
    t2 = doc.add_table(rows=3, cols=2)
    set_table_props(t2)
    
    fields_t2 = [
        ('課堂案例深度思辨\n【6 歲女孩的 25 歲 AI 照】\n（審美規訓與想像力殖民）', '【案例連結】：https://www.facebook.com/sungchien/posts/pfbid0pCZDZkDEKXkyTkoroqZsWEzNouWbFqEcJ1BK2WgcZfFKhMpJiEqvHrVREodG4VK2l\n【小組思辨記錄（約 60 字）】：\n1. 為何 AI 將 6 歲女孩親手剪貼的未來護理師自畫像，替換為「拉直波浪髮、精緻成人妝容」的標準化外貌？這對孩童自我認同與原創想像力有何傷害？\n__________________________________________________________________________________\n2. 學校未獲家長同意將學童臉孔上傳第三方 AI，對本組專題田野採集有何生物特徵隱私啟示？\n__________________________________________________________________________________'),
        ('預設提示詞偏誤誘發實驗\n（Bias Elicitation Test）', '【輸入之未約束籠統提示詞（Unconstrained Prompt）】：\nPrompt: 「________________________________________________________________________」\n【AI 預設產出之演算法偏誤觀察】：\n□ 性別與職業刻板印象  □ 歐美白人中心視角  □ 網紅審美濾鏡與年齡規訓  □ 在地文化失真\n具體偏誤現象描述（約 40 字）：_____________________________________________________'),
        ('偏誤成因診斷\n（統計多數霸權分析）', '【為什麼模型會給出上述預設結果？】：\n分析訓練資料集（商業圖庫、社群高讚照片）如何將統計多數放大為唯一標準，並抹殺多元群體表徵：\n__________________________________________________________________________________')
    ]
    
    for i, (label, val) in enumerate(fields_t2):
        c0 = t2.cell(i, 0)
        c1 = t2.cell(i, 1)
        set_cell_format(c0, width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER)
        set_cell_format(c1, width_dxa=CONTENT_WIDTH_DXA)
        
        p_lbl = c0.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(9.5)
        
        p_val = c1.paragraphs[0]
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.0)

    # ------------------ 表格三：SDGs 專題三大倫理地雷審計與去偏誤重構 ------------------
    add_heading_paragraph('【模組三：SDGs 專題三大倫理地雷自我審計與 S-S-L-M 去偏誤賦權重構】')
    
    t3 = doc.add_table(rows=3, cols=2)
    set_table_props(t3)
    
    fields_t3 = [
        ('三大倫理地雷自我排查\n（針對第 4 週海報與第 5 週文案）', '【逐項檢視本組海報構圖與倡議文案】：\n1. 【剝削式悲情（Poverty Porn）】：是否刻意渲染弱勢者狼狽哭泣、損害當事人尊嚴？ □ 無  □ 需修正\n2. 【救世主情結（Savior Complex）】：是否將倡議者描繪為高高在上的施捨英雄？ □ 無  □ 需修正\n3. 【視覺與文案漂綠（Greenwashing）】：是否捏造違背生態科學之虛假環保奇景？ □ 無  □ 需修正'),
        ('S-S-L-M 去偏誤與賦權重構\n（Inclusive & Empowering Prompt）', '【植入三大去偏誤法則後之升級版視覺提示詞 / 畫面設定（約 60 字）】：\n• 在地文化與多元表徵錨定：________________________________________________________\n• 保留真實自然肌理（Authentic Imperfection，拒絕塑膠濾鏡）：________________________\n• 平視賦權視角（Eye-level shot，展現受眾自主韌性與尊嚴）：__________________________'),
        ('去說教之共融文案微調\n（Inclusive Copywriting）', '【經倫理審計優化後之最終版 Campaign Slogan 與核心行動呼籲】：\n__________________________________________________________________________________')
    ]
    
    for i, (label, val) in enumerate(fields_t3):
        c0 = t3.cell(i, 0)
        c1 = t3.cell(i, 1)
        set_cell_format(c0, width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER)
        set_cell_format(c1, width_dxa=CONTENT_WIDTH_DXA)
        
        p_lbl = c0.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(9.5)
        
        p_val = c1.paragraphs[0]
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.0)

    # ------------------ 表格四：SynthID 防偽檢測、反向搜尋與透明揭露聲明 ------------------
    add_heading_paragraph('【模組四：SynthID 數位浮水印查核、反向圖片搜尋溯源與透明揭露聲明】')
    
    t4 = doc.add_table(rows=3, cols=2)
    set_table_props(t4)
    
    fields_t4 = [
        ('Google SynthID 浮水印\n與「關於這張圖片」實測', '【檢測本組第 4 週 Google AI 生成海報】：\n透過 Google 圖片搜尋「關於這張圖片（About this image）」或 SynthID 檢測器查驗結果：\n□ 成功檢出「由 Google AI 製作（SynthID 數位浮水印 / C2PA 內容憑證）」標記\n□ 因歷經大幅度後製圖層合成，系統顯示部分修改，本組將於作品旁主動標示 AI 協作說明'),
        ('外部素材與新聞圖片\n反向圖片搜尋溯源查核\n（Reverse Image Search）', '【查核專題參考或簡報引用之外部真實圖片／統計圖表】：\n• 查核圖片主題描述：______________________________________________________________\n• Google 首次索引出現年份：______ 年 ｜ 最原始權威發布機構：________________________\n• 溯源結論：□ 確認為真實紀實照片／官方圖表，絕無「舊圖當新災情」或「誤引 AI 假圖」'),
        ('專案《負責任 AI 創作與\n透明揭露聲明》設計\n（Responsible AI Disclosure）', '【撰寫將印製於第 10 週活動網頁頁尾（Footer）與期末發表簡報之正式揭露聲明（約 60–80 字）】：\n「【負責任 AI 創作與透明揭露聲明】：本專案（《____________________》）之視覺意象係由團隊主導原創構思，結合 Google AI 多模態工具協同生成並經人工編修（內含 SynthID 防偽標記）；全案引用之統計數據 100% 溯源於________________________真實文獻，絕無虛構事實或漂綠誇大。」')
    ]
    
    for i, (label, val) in enumerate(fields_t4):
        c0 = t4.cell(i, 0)
        c1 = t4.cell(i, 1)
        set_cell_format(c0, width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER)
        set_cell_format(c1, width_dxa=CONTENT_WIDTH_DXA)
        
        p_lbl = c0.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(9.5)
        
        p_val = c1.paragraphs[0]
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.0)

    # ------------------ 表格五：倫理紅隊同儕互評與教師審查 ------------------
    add_heading_paragraph('【模組五：課堂倫理紅隊突襲（Ethical Red Teaming）互評與教師審查簽章】')
    
    t5 = doc.add_table(rows=2, cols=2)
    set_table_props(t5)
    
    fields_t5 = [
        ('課堂 3 分鐘倫理紅隊互評\n（對手組別：第 _____ 組）', '【對手紅隊稽核員之挑刺焦點（指出本組海報/文案潛藏之刻板印象、說教感或漂綠風險）】：\n__________________________________________________________________________________\n【本組藍隊總監之防禦與重構回應（說明如何以平視賦權視角與透明揭露化解疑慮）】：\n__________________________________________________________________________________'),
        ('課堂即時諮詢與\n審查回饋欄\n（授課教師／助教專用）', '【專案倫理與防偽查核審查判定】：\n□ 審核通過（深刻反思演算法偏見、完成去偏誤賦權重構、SynthID 查核確實、透明聲明完備）\n□ 建議微調（視覺或文案仍帶有些微刻板印象或道德說教感，建議依紅隊意見微調）\n□ 需重新修正（存在消費弱勢尊嚴或誇大漂綠之倫理風險，需重構提示詞與論述）\n\n指導回饋與期中考前總評建議：\n__________________________________________________________________________________\n\n教師／助教簽章：______________________ ｜ 審核日期：2026 年 10 月 28 日')
    ]
    
    for i, (label, val) in enumerate(fields_t5):
        c0 = t5.cell(i, 0)
        c1 = t5.cell(i, 1)
        if i == 1:
            set_cell_format(c0, width_dxa=LABEL_WIDTH_DXA, fill_hex="FDEDEC")
            set_cell_format(c1, width_dxa=CONTENT_WIDTH_DXA, fill_hex="FDFEFE")
            p_lbl = c0.paragraphs[0]
            r_lbl = p_lbl.add_run(label)
            r_lbl.font.bold = True
            r_lbl.font.size = Pt(9.5)
            r_lbl.font.color.rgb = COLOR_REVIEW
        else:
            set_cell_format(c0, width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER)
            set_cell_format(c1, width_dxa=CONTENT_WIDTH_DXA)
            p_lbl = c0.paragraphs[0]
            r_lbl = p_lbl.add_run(label)
            r_lbl.font.bold = True
            r_lbl.font.size = Pt(9.5)
        
        p_val = c1.paragraphs[0]
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.0)

    # 存檔
    output_path = 'SDGs_專題生成內容倫理偏誤審計與SynthID防偽查核單.docx'
    doc.save(output_path)
    print("Successfully generated week 7 worksheet docx file.")

if __name__ == '__main__':
    create_document()
