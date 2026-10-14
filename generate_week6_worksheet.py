import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = docx.Document()
    
    # 頁面邊界設定 (2.032 cm / 0.8 inch 全等邊界)
    section = doc.sections[0]
    section.top_margin = Cm(2.032)
    section.bottom_margin = Cm(2.032)
    section.left_margin = Cm(2.032)
    section.right_margin = Cm(2.032)
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    
    # 全域樣式設定
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Microsoft JhengHei'
    normal_style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    normal_style.font.size = Pt(9.5)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(3)

    # 樣式常數定義
    COLOR_PRIMARY = RGBColor(0x1a, 0x56, 0xb8)     # 皇家藍 (標題與章節標題)
    COLOR_MUTED = RGBColor(0x7f, 0x8c, 0x8d)       # 灰字 (頂部引導備註)
    COLOR_BODY = RGBColor(0x55, 0x55, 0x55)        # 說明次要字
    COLOR_REVIEW = RGBColor(0xc0, 0x39, 0x2b)      # 緋紅色 (教師審查欄)
    
    FILL_HEADER = "F0F4F8"                         # 欄位標籤淺藍灰底色
    BORDER_COLOR = "D0D7DE"                        # 表格細邊框淺灰色
    
    LABEL_WIDTH_DXA = 2880                         # 5.08 cm
    CONTENT_WIDTH_DXA = 7056                       # 12.446 cm
    
    def set_cell_format(cell, width_dxa=None, fill_hex=None, v_align="center", pad_top=100, pad_bottom=100, pad_left=130, pad_right=130):
        tcPr = cell._tc.get_or_add_tcPr()
        
        # 寬度
        if width_dxa is not None:
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW')
                tcPr.append(tcW)
            tcW.set(qn('w:type'), 'dxa')
            tcW.set(qn('w:w'), str(width_dxa))
            
        # 垂直對齊
        vAlign = tcPr.find(qn('w:vAlign'))
        if vAlign is None:
            vAlign = OxmlElement('w:vAlign')
            tcPr.append(vAlign)
        vAlign.set(qn('w:val'), v_align)
        
        # 內距 padding
        tcMar = tcPr.find(qn('w:tcMar'))
        if tcMar is None:
            tcMar = OxmlElement('w:tcMar')
            tcPr.append(tcMar)
        for side, pad in [('top', pad_top), ('bottom', pad_bottom), ('left', pad_left), ('right', pad_right)]:
            m = tcMar.find(qn(f'w:{side}'))
            if m is None:
                m = OxmlElement(f'w:{side}')
                tcMar.append(m)
            m.set(qn('w:w'), str(pad))
            m.set(qn('w:type'), 'dxa')
            
        # 邊框
        borders_xml = f'''
        <w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>
        </w:tcBorders>
        '''
        existing_borders = tcPr.find(qn('w:tcBorders'))
        if existing_borders is not None:
            tcPr.remove(existing_borders)
        tcPr.append(parse_xml(borders_xml))
        
        # 底色
        if fill_hex:
            shd = tcPr.find(qn('w:shd'))
            if shd is None:
                shd = OxmlElement('w:shd')
                tcPr.append(shd)
            shd.set(qn('w:fill'), fill_hex)

    def set_table_props(table):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = table._tbl.tblPr
        tblW = tblPr.find(qn('w:tblW'))
        if tblW is not None:
            tblW.set(qn('w:type'), 'auto')
            tblW.set(qn('w:w'), '0')
        for row in table.rows:
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml('<w:cantSplit xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))

    def add_heading_paragraph(text, color=COLOR_PRIMARY, size_pt=11.5, space_before=10, space_after=3):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Microsoft JhengHei'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
        run.font.size = Pt(size_pt)
        run.font.bold = True
        run.font.color.rgb = color
        return p

    # ------------------ 文件頂部抬頭 ------------------
    p0 = doc.add_paragraph()
    p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p0.paragraph_format.space_after = Pt(2)
    r0 = p0.add_run('數位內容設計（Digital Content Design）｜ PBL 專題實作導引（第六週）')
    r0.font.name = 'Microsoft JhengHei'
    r0._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    r0.font.size = Pt(9.0)
    r0.font.color.rgb = COLOR_MUTED

    p1 = doc.add_paragraph()
    p1.paragraph_format.space_after = Pt(4)
    p1.paragraph_format.space_before = Pt(0)
    r1 = p1.add_run('第六週：SDGs 專題生成內容智財合規查核與模擬法庭攻防單')
    r1.font.name = 'Microsoft JhengHei'
    r1._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    r1.font.size = Pt(17.0)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_PRIMARY

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(8)
    p2.paragraph_format.line_spacing = 1.25
    r2 = p2.add_run(
        '本實作單用於第六週課堂工作坊（60 分鐘）。請各組調閱第四週產出之「Imagen 3 核心主視覺海報」與第五週收斂之「NotebookLM 數據錨點與 Campaign Slogan」，'
        '全面執行專案智慧財產權與法律合規審查：'
        '首先進行視覺要素與文案侵權自我審查，排查特定 IP 實質相似、建築商標與肖像權風險；'
        '接著運用 Google Gemini 開啟 Search Grounding（即時網路搜尋接地），檢索 2023 至 2026 年最新台灣智慧財產法院判決或智財局官方函釋，錨定客觀法律依據；'
        '隨後依據著作權法合理使用（Fair Use）四大準則，建構小組之「轉化性使用（Transformative Use）」法律防禦論述，並記錄人類創作性控制（Human Creative Control）歷程；'
        '最後參與課堂 3 分鐘模擬法庭微辯論，直面原告指控並展開專業答辯。'
        '本表單經教師/助教課堂查核簽章後，將作為期末專案企劃書之正式智財合規章節依據。'
    )
    r2.font.name = 'Microsoft JhengHei'
    r2._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    r2.font.size = Pt(9.0)
    r2.font.color.rgb = COLOR_BODY

    # ------------------ 表格一：基本資訊與既有資產 ------------------
    add_heading_paragraph('【模組一：專案基本資訊與既有產出資產回顧】')
    
    t1 = doc.add_table(rows=4, cols=2)
    set_table_props(t1)
    
    fields_t1 = [
        ('小組組別與專案名稱', '組別：第 ______ 組 ｜ 專案名稱：《__________________________________________________》'),
        ('選定之 SDGs 目標', '第 ______ 號核心目標：________________________（中文完整名稱）'),
        ('第四週與第五週既有資產\n（主視覺與 Slogan 回顧）', '【第四週主視覺海報核心特徵】：\n__________________________________________________________________________________\n【第五週 Campaign 一句話 Slogan】：\n《________________________________________________________________》'),
        ('小組成員與審查分工', '專案法務主審（PM）：______________ ｜ 視覺侵權排查（視覺）：______________\n文案商標排查（文案）：______________ ｜ 模擬法庭辯護代表：______________')
    ]
    
    for i, (label, val) in enumerate(fields_t1):
        c0 = t1.cell(i, 0)
        c1 = t1.cell(i, 1)
        set_cell_format(c0, width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER)
        set_cell_format(c1, width_dxa=CONTENT_WIDTH_DXA)
        
        p_lbl = c0.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(9.5)
        
        p_val = c1.paragraphs[0]
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.5)

    # ------------------ 表格二：視覺與文案侵權自我排查 ------------------
    add_heading_paragraph('【模組二：專案視覺要素與文案智財合規自我排查表】')
    
    t2 = doc.add_table(rows=3, cols=2)
    set_table_props(t2)
    
    fields_t2 = [
        ('海報主體角色與肖像權\n（實質相似度自我檢驗）', '【主體角色外觀比對】：是否有任何輪廓、髮型、配色或配件，與既有動漫、公仔或遊戲知名 IP 產生混淆或實質相似？\n□ 無侵權疑慮（純原創特徵）  □ 有特定知名風格借鑑但無具體角色抄襲  □ 有高風險雷同（需修改）\n具體審查說明（約 40 字）：_____________________________________________________'),
        ('背景建築、商標與標誌\n（商標法與附隨物檢視）', '【畫面背景與物件排查】：畫面中是否有清晰可辨識之知名商業品牌商標（如 Logo、鞋服標誌）或受保護特定現代建築著作？\n□ 無任何商業標誌  □ 含有部分現實地標但屬於街景合理附隨  □ 誤入商業商標（需 Inpainting 消除）\n具體排查結論：_________________________________________________________________'),
        ('文案數據與 Campaign 標語\n（商標撞車與著作權排查）', '【Slogan 商標查核與數據引述】：\n1. Slogan 是否與市面上已註冊之著名商標字串撞車？（已於智慧財產局商標檢索系統查證）：□ 是  □ 否\n2. 第五週文獻量化數據是否均已標記機構出處與出版年份？：□ 均已完整標記  □ 尚需補齊出處')
    ]
    
    for i, (label, val) in enumerate(fields_t2):
        c0 = t2.cell(i, 0)
        c1 = t2.cell(i, 1)
        set_cell_format(c0, width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER)
        set_cell_format(c1, width_dxa=CONTENT_WIDTH_DXA)
        
        p_lbl = c0.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(9.5)
        
        p_val = c1.paragraphs[0]
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.0)

    # ------------------ 表格三：Gemini + Search Grounding 判例檢索 ------------------
    add_heading_paragraph('【模組三：Gemini + Search Grounding 判例檢索與司法裁判溯源】')
    
    t3 = doc.add_table(rows=3, cols=2)
    set_table_props(t3)
    
    fields_t3 = [
        ('CRISP 檢索提示工程紀錄\n（法律諮詢 Agent 設定）', '【在 Gemini 中輸入之 Context 與 Role 設定】：\n「你是一位專精於台灣《著作權法》與生成式 AI 智財訴訟的資深法務顧問……」\n【核心檢索任務 Instruction 與檢索關鍵字】：\n__________________________________________________________________________________'),
        ('查得之法院裁判/主管機關函釋\n（真實案號與裁決要旨）', '【裁判字號 / 行政函釋字號】：\n案號：【________________________________________】（裁判/發文年份：______ 年）\n【法院或主管機關核心裁判見解摘要（約 50 字）】：\n__________________________________________________________________________________\n__________________________________________________________________________________'),
        ('Grounded 即時來源查證\n（點擊官方超連結核對）', '【正式官方原文來源網址（可點擊之司法院裁判書系統或智財局網址）】：\nURL: https://_____________________________________________________________________\n【核驗確認】：□ 已親自點擊超連結，核對判決書主文與當事人無誤，確認非 AI 虛構幻覺！')
    ]
    
    for i, (label, val) in enumerate(fields_t3):
        c0 = t3.cell(i, 0)
        c1 = t3.cell(i, 1)
        set_cell_format(c0, width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER)
        set_cell_format(c1, width_dxa=CONTENT_WIDTH_DXA)
        
        p_lbl = c0.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(9.5)
        
        p_val = c1.paragraphs[0]
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.0)

    # ------------------ 表格四：商業合理使用防禦答辯書 ------------------
    add_heading_paragraph('【模組四：商業合理使用（Fair Use）四大準則專案防禦答辯書】')
    
    t4 = doc.add_table(rows=4, cols=2)
    set_table_props(t4)
    
    fields_t4 = [
        ('準則一：利用之目的與性質\n【轉化性使用深度】\n（決定勝敗的核心要件）', '【轉化性使用（Transformative Use）論述（約 60 字）】：\n說明本專案如何超越原著作的既有功能，賦予全新的社會永續意義、批判洞察與審美價值：\n__________________________________________________________________________________\n__________________________________________________________________________________'),
        ('準則二：被利用著作之性質\n（事實性 vs. 高度原創性）', '【論述引述素材之性質】：\n本專案引用者主要為：□ 客觀事實、統計數據與官方白皮書（受保護程度較低，合理使用空間大）\n□ 部份藝術風格與構圖靈感（屬思想與風格層次，非具體表達抄襲）\n說明：_________________________________________________________________________'),
        ('準則三：所利用質量與比例\n（非核心實質部分之證明）', '【數量與質量核心論述】：\n本專案利用他人素材之篇幅僅佔原著作極微小比例，且未掠取原著作最具商業價值之「核心心臟（Heart of the Work）」：\n證明說明：_____________________________________________________________________'),
        ('準則四：對潛在市場之影響\n（無市場替代效應論證）', '【市場影響評估】：\n本專案產出是否會取代原著作之市場銷售或訂閱？\n□ 絕無替代效應；本專案為推動 SDGs 倡議，反而具有擴大公眾對該議題關注與原報告閱讀之正面互補效應。')
    ]
    
    for i, (label, val) in enumerate(fields_t4):
        c0 = t4.cell(i, 0)
        c1 = t4.cell(i, 1)
        set_cell_format(c0, width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER)
        set_cell_format(c1, width_dxa=CONTENT_WIDTH_DXA)
        
        p_lbl = c0.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(9.5)
        
        p_val = c1.paragraphs[0]
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.0)

    # ------------------ 表格五：人類創作性控制歷程證明 ------------------
    add_heading_paragraph('【模組五：人類創作性控制歷程證明（Human Authorship Log）】')
    
    t5 = doc.add_table(rows=2, cols=2)
    set_table_props(t5)
    
    fields_t5 = [
        ('提示工程與參數迭代歷程\n（舉證非一鍵隨機產物）', '【結構化 Prompt 調校歷程摘要】：\n共歷經 ______ 輪提示詞微調，調校參數包含 Temperature、Seed、CFG 權重與 S-S-L-M 屬性，\n從 ______ 張候選生成影像中進行美學挑選、過濾與淘汰，展現明確的人類審美意圖。'),
        ('後期人工實質編修與合成\n（主張著作權保護之核心基石）', '【後期人工設計勞力投入紀錄（約 60 字）】：\n（例如：使用 Photoshop / Figma 進行局部 Inpainting 消除瑕疵、重新調色、向量標題字型設計排版、圖層多重合成）：\n__________________________________________________________________________________\n【工程檔案封存證明】：□ 已將 PSD/AI/Figma 原始分層工作檔備份於小組雲端，作為日後原創性舉證資產。')
    ]
    
    for i, (label, val) in enumerate(fields_t5):
        c0 = t5.cell(i, 0)
        c1 = t5.cell(i, 1)
        set_cell_format(c0, width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER)
        set_cell_format(c1, width_dxa=CONTENT_WIDTH_DXA)
        
        p_lbl = c0.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(9.5)
        
        p_val = c1.paragraphs[0]
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.0)

    # ------------------ 表格六：模擬法庭微辯論與教師審查 ------------------
    add_heading_paragraph('【模組六：課堂模擬法庭微辯論攻防紀錄與教師審查簽章】')
    
    t6 = doc.add_table(rows=2, cols=2)
    set_table_props(t6)
    
    fields_t6 = [
        ('模擬法庭 3 分鐘攻防紀錄\n（對手組別：第 _____ 組）', '【對手原告律師之攻擊焦點（指控本組侵權之核心理由）】：\n__________________________________________________________________________________\n【本組被告團隊之防禦答辯（依據合理使用與判例之抗辯主張）】：\n__________________________________________________________________________________'),
        ('課堂即時諮詢與\n審查回饋欄\n（授課教師／助教專用）', '【專案智財合規審查判定】：\n□ 合規通過（法律概念清晰、判例溯源真實、轉化性論述扎實、創作歷程完備）\n□ 建議微調（特定視覺特徵或文案仍有輕微侵權疑慮，建議進行局部修整）\n□ 重新重繪（存在重大實質相似或商標侵權高風險，強制更換提示詞重製）\n\n指導回饋與具體修改意見：\n__________________________________________________________________________________\n\n教師／助教簽章：______________________ ｜ 審核日期：2026 年 10 月 21 日')
    ]
    
    for i, (label, val) in enumerate(fields_t6):
        c0 = t6.cell(i, 0)
        c1 = t6.cell(i, 1)
        if i == 1:
            set_cell_format(c0, width_dxa=LABEL_WIDTH_DXA, fill_hex="FDEDEC")
            set_cell_format(c1, width_dxa=CONTENT_WIDTH_DXA, fill_hex="FDFEFE")
            p_lbl = c0.paragraphs[0]
            r_lbl = p_lbl.add_run(label)
            r_lbl.font.bold = True
            r_lbl.font.size = Pt(9.5)
            r_lbl.font.color.rgb = COLOR_REVIEW
        else:
            set_cell_format(c0, width_dxa=LABEL_WIDTH_DXA, fill_hex=FILL_HEADER)
            set_cell_format(c1, width_dxa=CONTENT_WIDTH_DXA)
            p_lbl = c0.paragraphs[0]
            r_lbl = p_lbl.add_run(label)
            r_lbl.font.bold = True
            r_lbl.font.size = Pt(9.5)
        
        p_val = c1.paragraphs[0]
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.0)

    # 存檔
    output_path = 'SDGs_專題生成內容智財合規查核與模擬法庭攻防單.docx'
    doc.save(output_path)
    print("Successfully generated week 6 worksheet docx file.")

if __name__ == '__main__':
    create_document()
